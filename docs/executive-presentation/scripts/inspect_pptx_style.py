from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")

base = Path(r"c:\Users\mosmo\OneDrive\Desktop\it-portfolio-system\docs\executive-presentation")
NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}


def dump_theme(prs: Presentation) -> None:
    # theme lives in slide master part
    for mi, master in enumerate(prs.slide_masters):
        part = master.part
        # related theme
        for rel in part.rels.values():
            if "theme" in rel.reltype:
                theme = etree.fromstring(rel.target_part.blob)
                print(f"\nTHEME master[{mi}] {rel.target_ref}")
                # color scheme
                scheme = theme.find(".//a:clrScheme", NS)
                if scheme is not None:
                    print("  clrScheme:", scheme.get("name"))
                    for child in scheme:
                        tag = etree.QName(child).localname
                        srgb = child.find(".//a:srgbClr", NS)
                        sysClr = child.find(".//a:sysClr", NS)
                        if srgb is not None:
                            print(f"    {tag}: #{srgb.get('val')}")
                        elif sysClr is not None:
                            print(f"    {tag}: sys={sysClr.get('val')} last={sysClr.get('lastClr')}")
                # fonts
                major = theme.find(".//a:majorFont/a:latin", NS)
                minor = theme.find(".//a:minorFont/a:latin", NS)
                major_ea = theme.find(".//a:majorFont/a:ea", NS)
                minor_ea = theme.find(".//a:minorFont/a:ea", NS)
                major_cs = theme.find(".//a:majorFont/a:cs", NS)
                minor_cs = theme.find(".//a:minorFont/a:cs", NS)
                print("  major latin:", major.get("typeface") if major is not None else None)
                print("  minor latin:", minor.get("typeface") if minor is not None else None)
                print("  major ea:", major_ea.get("typeface") if major_ea is not None else None)
                print("  minor ea:", minor_ea.get("typeface") if minor_ea is not None else None)
                print("  major cs:", major_cs.get("typeface") if major_cs is not None else None)
                print("  minor cs:", minor_cs.get("typeface") if minor_cs is not None else None)


def inspect_slide(prs: Presentation, idx: int) -> None:
    slide = prs.slides[idx]
    print(f"\n===== SLIDE {idx+1} shapes={len(slide.shapes)} =====")
    # background
    bg = slide.background
    try:
        fill = bg.fill
        print("background fill type:", fill.type)
        try:
            print("  bg rgb:", fill.fore_color.rgb)
        except Exception:
            pass
        try:
            print("  bg theme:", fill.fore_color.theme_color)
        except Exception:
            pass
    except Exception as e:
        print("bg err", e)

    # layout name
    try:
        print("layout:", slide.slide_layout.name)
    except Exception:
        pass

    for sh in slide.shapes:
        left = round(sh.left.inches, 2)
        top = round(sh.top.inches, 2)
        w = round(sh.width.inches, 2)
        h = round(sh.height.inches, 2)
        info = f"{sh.shape_type} {sh.name!r} @({left},{top}) {w}x{h}"
        try:
            if hasattr(sh, "fill") and sh.fill.type is not None:
                info += f" fill={sh.fill.type}"
                try:
                    info += f":{sh.fill.fore_color.rgb}"
                except Exception:
                    try:
                        info += f":theme={sh.fill.fore_color.theme_color}"
                    except Exception:
                        pass
        except Exception:
            pass
        print(info)
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            print(f"  PICTURE {sh.image.content_type} {len(sh.image.blob)} bytes")
        if sh.has_text_frame:
            for pi, p in enumerate(sh.text_frame.paragraphs):
                if not p.text.strip():
                    continue
                runs_meta = []
                for r in p.runs:
                    bit = {"t": r.text[:40], "font": r.font.name, "sz": r.font.size.pt if r.font.size else None}
                    try:
                        bit["rgb"] = str(r.font.color.rgb)
                    except Exception:
                        try:
                            bit["theme"] = str(r.font.color.theme_color)
                        except Exception:
                            pass
                    runs_meta.append(bit)
                print(f"  P{pi}: {p.text.strip()[:100]!r} -> {runs_meta}")


def main() -> None:
    style = Presentation(str(base / "style.pptx"))
    print("STYLE slides", len(style.slides), "layouts", [l.name for l in style.slide_layouts])
    dump_theme(style)
    for i in range(len(style.slides)):
        inspect_slide(style, i)

    print("\n\n################ CENTRIX ################")
    cx = Presentation(str(base / "CENTRIX.pptx"))
    print("CENTRIX slides", len(cx.slides))
    dump_theme(cx)
    for i in range(min(3, len(cx.slides))):
        inspect_slide(cx, i)
    print("\nAll CENTRIX slide titles/first texts:")
    for i, slide in enumerate(cx.slides, 1):
        texts = []
        for sh in slide.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    if p.text.strip():
                        texts.append(p.text.strip())
                        break
            if len(texts) >= 3:
                break
        print(f"  {i:02d}: {texts}")


if __name__ == "__main__":
    main()
