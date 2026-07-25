from __future__ import annotations

import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")

base = Path(r"c:\Users\mosmo\OneDrive\Desktop\it-portfolio-system\docs\executive-presentation")
out = base / "assets" / "style-extracted"
out.mkdir(parents=True, exist_ok=True)

prs = Presentation(str(base / "style.pptx"))
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}

print("LAYOUT DETAILS")
for layout in prs.slide_layouts:
    name = layout.name
    print(f"\nLAYOUT: {name}")
    # background
    bg = layout.background
    try:
        print("  fill type", bg.fill.type)
    except Exception as e:
        print("  fill err", e)

    # shapes on layout
    for sh in layout.shapes:
        print(f"  shape {sh.shape_type} {sh.name!r} ({sh.left.inches:.2f},{sh.top.inches:.2f}) {sh.width.inches:.2f}x{sh.height.inches:.2f}")
        if sh.has_text_frame:
            t = " | ".join(p.text.strip() for p in sh.text_frame.paragraphs if p.text.strip())[:80]
            if t:
                print(f"    text: {t}")

    # extract images from layout part relationships
    part = layout.part
    for rel in part.rels.values():
        if "image" in rel.reltype:
            blob = rel.target_part.blob
            ct = rel.target_part.content_type
            ext = ".png" if "png" in ct else ".jpg" if "jpeg" in ct or "jpg" in ct else ".bin"
            fname = out / f"layout_{name.replace(' ', '_')}{ext}"
            # uniquify
            i = 1
            while fname.exists():
                fname = out / f"layout_{name.replace(' ', '_')}_{i}{ext}"
                i += 1
            fname.write_bytes(blob)
            print(f"  IMAGE saved {fname.name} ({len(blob)} bytes, {ct})")

# also extract slide 4 pictures as full backgrounds
slide4 = prs.slides[3]
for idx, sh in enumerate(slide4.shapes):
    if sh.shape_type is not None and sh.shape_type == 13:  # PICTURE
        blob = sh.image.blob
        ct = sh.image.content_type
        ext = ".png" if "png" in ct else ".jpg"
        fname = out / f"slide4_pic_{idx}{ext}"
        fname.write_bytes(blob)
        print(f"SLIDE4 pic {fname.name} @({sh.left.inches:.2f},{sh.top.inches:.2f}) {sh.width.inches:.2f}x{sh.height.inches:.2f} {len(blob)}b")

# master images
for mi, master in enumerate(prs.slide_masters):
    for rel in master.part.rels.values():
        if "image" in rel.reltype:
            blob = rel.target_part.blob
            ct = rel.target_part.content_type
            ext = ".png" if "png" in ct else ".jpg"
            fname = out / f"master{mi}_{Path(rel.target_ref).name}"
            if not fname.suffix:
                fname = fname.with_suffix(ext)
            fname.write_bytes(blob)
            print(f"MASTER IMAGE {fname.name} {len(blob)}b")

print("\nExtracted to", out)
for p in sorted(out.glob("*")):
    print(" ", p.name, p.stat().st_size)
