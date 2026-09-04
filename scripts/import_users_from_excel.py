#!/usr/bin/env python3
"""Import users from the Users Excel sheet into the Laravel /api/v1/users API.

The sheet in users.png has headers on row 2, columns B–F:

    Sl. No. | Display Name | Email | Applications | Phone Number

Display Name is split on the first space into first_name / last_name.
Phone numbers are normalized so they always begin with +966.

Examples:
    python scripts/import_users_from_excel.py --excel users.xlsx --dry-run
    python scripts/import_users_from_excel.py --excel users.xlsx --apply
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import time
import urllib.error
import urllib.request
from dataclasses import dataclass
from pathlib import Path
from typing import Any

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from excel_import.normalize import (  # noqa: E402
    cell_text,
    header_key,
    is_blank,
    is_valid_email,
    normalize_email,
    normalize_saudi_phone,
    normalize_whitespace,
    split_display_name,
)

try:
    from openpyxl import load_workbook
    from openpyxl.worksheet.worksheet import Worksheet
except ImportError as exc:  # pragma: no cover
    raise SystemExit(
        "openpyxl is required. Install it with:\n"
        "  pip install -r scripts/requirements-migration.txt"
    ) from exc


DISPLAY_NAME_HEADERS = {
    "display name",
    "name",
    "employee name",
    "full name",
    "fullname",
}
EMAIL_HEADERS = {"email", "e-mail", "email address", "mail"}
PHONE_HEADERS = {
    "phone number",
    "phone",
    "mobile",
    "mobile number",
    "mobile no",
    "mobile no.",
    "cellphone",
    "cell phone",
}

DEFAULT_API_URL = os.environ.get("API_BASE_URL", "http://127.0.0.1:8000/api/v1")
DEFAULT_AUTH_EMAIL = os.environ.get("AUTH_EMAIL", "admin@zatca.gov.sa")
DEFAULT_AUTH_PASSWORD = os.environ.get("AUTH_PASSWORD", "password")
DEFAULT_USER_PASSWORD = os.environ.get("USER_DEFAULT_PASSWORD", "password")


@dataclass(frozen=True)
class ColumnMap:
    sheet_name: str
    header_row: int
    display_name: int
    email: int
    phone: int


@dataclass
class PreparedUser:
    row: int
    display_name: str
    first_name: str
    last_name: str
    email: str
    phone: str
    skipped_reason: str | None = None

    def payload(self, password: str, roles: list[str]) -> dict[str, Any]:
        body: dict[str, Any] = {
            "first_name": self.first_name,
            "last_name": self.last_name,
            "email": self.email,
            "password": password,
            "password_confirmation": password,
            "phone": self.phone,
            "is_active": True,
        }
        if roles:
            body["roles"] = roles
        return body


class ApiError(RuntimeError):
    def __init__(self, status: int, message: str, errors: Any = None) -> None:
        super().__init__(message)
        self.status = status
        self.errors = errors


class UsersApiClient:
    def __init__(self, base_url: str, timeout: int = 30) -> None:
        self.base_url = base_url.rstrip("/")
        self.timeout = timeout
        self.token: str | None = None

    def login(self, email: str, password: str) -> None:
        data = self._request(
            "POST",
            "/auth/login",
            {"email": email, "password": password},
            authenticated=False,
        )
        token = (data or {}).get("token")
        if not token:
            raise ApiError(0, "Login succeeded but no token was returned.")
        self.token = str(token)

    def existing_emails(self) -> set[str]:
        emails: set[str] = set()
        page = 1
        last_page = 1
        while page <= last_page:
            data = self._request("GET", f"/users?page={page}&per_page=100")
            items = (data or {}).get("items") or []
            pagination = (data or {}).get("pagination") or {}
            last_page = int(pagination.get("last_page") or 1)
            for item in items:
                email = normalize_email(str(item.get("email") or ""))
                if email:
                    emails.add(email)
            page += 1
        return emails

    def create_user(self, payload: dict[str, Any]) -> dict[str, Any]:
        return self._request("POST", "/users", payload) or {}

    def _request(
        self,
        method: str,
        path: str,
        body: dict[str, Any] | None = None,
        authenticated: bool = True,
    ) -> dict[str, Any] | None:
        url = f"{self.base_url}{path}"
        headers = {
            "Accept": "application/json",
            "Content-Type": "application/json",
        }
        if authenticated:
            if not self.token:
                raise ApiError(401, "Not authenticated. Login first.")
            headers["Authorization"] = f"Bearer {self.token}"

        encoded = None if body is None else json.dumps(body).encode("utf-8")
        request = urllib.request.Request(url, data=encoded, headers=headers, method=method)
        try:
            with urllib.request.urlopen(request, timeout=self.timeout) as response:
                raw = response.read().decode("utf-8")
        except urllib.error.HTTPError as exc:
            raw = exc.read().decode("utf-8", errors="replace")
            payload = _decode_json(raw)
            message = str(payload.get("message") or exc.reason or "API request failed")
            raise ApiError(exc.code, message, payload.get("errors")) from exc
        except urllib.error.URLError as exc:
            raise ApiError(0, f"Could not reach API at {url}: {exc.reason}") from exc

        payload = _decode_json(raw)
        if payload.get("success") is False:
            raise ApiError(
                0,
                str(payload.get("message") or "API request failed"),
                payload.get("errors"),
            )
        return payload.get("data")


def _decode_json(raw: str) -> dict[str, Any]:
    if not raw.strip():
        return {}
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        return {"message": raw}
    if isinstance(parsed, dict):
        return parsed
    return {"data": parsed}


def detect_users_columns(sheet: Worksheet, header_row: int | None = None) -> ColumnMap:
    start = header_row or 1
    end = header_row or min(sheet.max_row or 1, 15)

    best: ColumnMap | None = None
    best_score = -1
    for row_idx in range(start, end + 1):
        display_col = email_col = phone_col = None
        for col_idx in range(1, min((sheet.max_column or 1), 30) + 1):
            key = header_key(sheet.cell(row_idx, col_idx).value)
            if key in DISPLAY_NAME_HEADERS and display_col is None:
                display_col = col_idx
            elif key in EMAIL_HEADERS and email_col is None:
                email_col = col_idx
            elif key in PHONE_HEADERS and phone_col is None:
                phone_col = col_idx
        score = sum(col is not None for col in (display_col, email_col, phone_col))
        if score == 3 and score > best_score:
            best_score = score
            best = ColumnMap(sheet.title, row_idx, display_col, email_col, phone_col)

    if best is None:
        raise RuntimeError(
            f"Sheet '{sheet.title}' is missing Display Name, Email, or Phone Number headers."
        )
    return best


def find_users_sheet(workbook, sheet_name: str | None = None) -> Worksheet:
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise RuntimeError(
                f"Sheet '{sheet_name}' was not found. Available: {workbook.sheetnames}"
            )
        return workbook[sheet_name]

    for name in workbook.sheetnames:
        if name.strip().casefold() in {"users", "user", "employees"}:
            return workbook[name]

    for name in workbook.sheetnames:
        sheet = workbook[name]
        try:
            detect_users_columns(sheet)
            return sheet
        except RuntimeError:
            continue

    raise RuntimeError(
        "Could not detect a Users sheet. Expected headers: Display Name, Email, Phone Number."
    )


def extract_users(
    path: Path,
    sheet_name: str | None = None,
    header_row: int | None = None,
    limit: int | None = None,
) -> tuple[list[PreparedUser], ColumnMap]:
    workbook = load_workbook(filename=path, data_only=True, read_only=True)
    try:
        sheet = find_users_sheet(workbook, sheet_name)
        columns = detect_users_columns(sheet, header_row)
        users: list[PreparedUser] = []
        seen_emails: set[str] = set()
        max_row = sheet.max_row or columns.header_row

        for row_idx in range(columns.header_row + 1, max_row + 1):
            display_name = normalize_whitespace(
                cell_text(sheet.cell(row_idx, columns.display_name).value)
            )
            email_raw = cell_text(sheet.cell(row_idx, columns.email).value)
            phone_raw = sheet.cell(row_idx, columns.phone).value

            if is_blank(display_name) and is_blank(email_raw) and is_blank(phone_raw):
                continue

            first_name, last_name = split_display_name(display_name)
            email = normalize_email(email_raw)
            phone = normalize_saudi_phone(phone_raw)

            skipped_reason: str | None = None
            if not display_name or not first_name or not last_name:
                skipped_reason = "missing display name"
            elif not email:
                skipped_reason = "missing email"
            elif not is_valid_email(email):
                skipped_reason = f"invalid email '{email_raw}'"
            elif email in seen_emails:
                skipped_reason = "duplicate email in sheet"
            elif phone is None:
                skipped_reason = "missing or invalid phone number"

            if email:
                seen_emails.add(email)

            users.append(
                PreparedUser(
                    row=row_idx,
                    display_name=display_name,
                    first_name=first_name,
                    last_name=last_name,
                    email=email,
                    phone=phone or "",
                    skipped_reason=skipped_reason,
                )
            )
            if limit is not None and len(users) >= limit:
                break

        return users, columns
    finally:
        workbook.close()


def format_errors(errors: Any) -> str:
    if not errors:
        return ""
    if isinstance(errors, dict):
        parts: list[str] = []
        for field, messages in errors.items():
            if isinstance(messages, list):
                parts.append(f"{field}: {'; '.join(str(item) for item in messages)}")
            else:
                parts.append(f"{field}: {messages}")
        return " | ".join(parts)
    return str(errors)


def print_preview(users: list[PreparedUser], columns: ColumnMap) -> None:
    ready = [user for user in users if user.skipped_reason is None]
    skipped = [user for user in users if user.skipped_reason is not None]
    print(
        f"Sheet '{columns.sheet_name}' header row {columns.header_row}: "
        f"{len(ready)} ready, {len(skipped)} skipped."
    )
    print()
    print(f"{'Row':<6} {'First name':<18} {'Last name':<22} {'Email':<36} {'Phone':<16} Status")
    print("-" * 120)
    for user in users:
        status = "OK" if user.skipped_reason is None else f"SKIP ({user.skipped_reason})"
        print(
            f"{user.row:<6} {user.first_name:<18} {user.last_name:<22} "
            f"{user.email:<36} {user.phone:<16} {status}"
        )


def import_users(
    users: list[PreparedUser],
    client: UsersApiClient,
    password: str,
    roles: list[str],
    delay: float,
    skip_existing: bool,
) -> tuple[int, int, int]:
    created = skipped = failed = 0
    existing: set[str] = set()
    if skip_existing:
        print("Loading existing users from API...")
        existing = client.existing_emails()
        print(f"Found {len(existing)} existing emails.")

    ready = [user for user in users if user.skipped_reason is None]
    for index, user in enumerate(ready, start=1):
        if skip_existing and user.email in existing:
            print(f"[{index}/{len(ready)}] skip existing {user.email}")
            skipped += 1
            continue
        try:
            created_user = client.create_user(user.payload(password, roles))
            user_id = created_user.get("id", "?")
            print(
                f"[{index}/{len(ready)}] created #{user_id} "
                f"{user.first_name} {user.last_name} <{user.email}> {user.phone}"
            )
            created += 1
            existing.add(user.email)
        except ApiError as exc:
            failed += 1
            details = format_errors(exc.errors)
            suffix = f" ({details})" if details else ""
            print(
                f"[{index}/{len(ready)}] FAILED {user.email}: "
                f"HTTP {exc.status} {exc}{suffix}"
            )
        if delay > 0:
            time.sleep(delay)

    sheet_skipped = sum(1 for user in users if user.skipped_reason is not None)
    return created, skipped + sheet_skipped, failed


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Extract users from the Excel Users sheet and insert them through "
            "POST /api/v1/users."
        )
    )
    parser.add_argument("--excel", required=True, type=Path, help="Path to the Excel workbook.")
    parser.add_argument("--sheet", default=None, help="Sheet name. Auto-detected when omitted.")
    parser.add_argument(
        "--header-row",
        type=int,
        default=None,
        help="1-based header row. Auto-detected (row 2 in the screenshot) when omitted.",
    )
    parser.add_argument("--api-url", default=DEFAULT_API_URL, help=f"API base URL. Default: {DEFAULT_API_URL}")
    parser.add_argument("--auth-email", default=DEFAULT_AUTH_EMAIL, help="Account with users.create permission.")
    parser.add_argument("--auth-password", default=DEFAULT_AUTH_PASSWORD, help="Login password for --auth-email.")
    parser.add_argument(
        "--user-password",
        default=DEFAULT_USER_PASSWORD,
        help="Password assigned to every imported user (min 8 characters).",
    )
    parser.add_argument(
        "--role",
        action="append",
        default=[],
        help="Spatie role to assign (repeatable), e.g. --role viewer",
    )
    parser.add_argument("--limit", type=int, default=None, help="Only process the first N data rows.")
    parser.add_argument("--delay", type=float, default=0.15, help="Seconds to wait between create calls.")
    parser.add_argument(
        "--apply",
        action="store_true",
        help="Create users through the API. Without this flag the script only prints a dry run.",
    )
    parser.add_argument(
        "--no-skip-existing",
        action="store_true",
        help="Do not pre-load existing emails; still fails on unique-email validation errors.",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    excel_path = args.excel.expanduser().resolve()
    if not excel_path.is_file():
        print(f"Excel file not found: {excel_path}", file=sys.stderr)
        return 1
    if len(args.user_password) < 8:
        print("Imported user password must be at least 8 characters.", file=sys.stderr)
        return 1

    users, columns = extract_users(
        excel_path,
        sheet_name=args.sheet,
        header_row=args.header_row,
        limit=args.limit,
    )
    print_preview(users, columns)

    ready = [user for user in users if user.skipped_reason is None]
    if not args.apply:
        print()
        print(f"Dry run only. Re-run with --apply to create {len(ready)} users via {args.api_url}/users")
        return 0

    if not ready:
        print("Nothing to import.")
        return 0

    client = UsersApiClient(args.api_url)
    print()
    print(f"Logging in as {args.auth_email}...")
    try:
        client.login(args.auth_email, args.auth_password)
    except ApiError as exc:
        print(f"Login failed: {exc}", file=sys.stderr)
        return 1

    created, skipped, failed = import_users(
        users,
        client,
        password=args.user_password,
        roles=args.role,
        delay=max(0.0, args.delay),
        skip_existing=not args.no_skip_existing,
    )
    print()
    print(f"Done. created={created} skipped={skipped} failed={failed}")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
