#!/usr/bin/env python3
"""Import applications from the PHASE2 Excel sheet into POST /api/v1/applications.

PHASE2 columns (header typically on row 1, B–H and optional extra columns):

    # | App Names | Technical Category | App Description | Under Operation?
    | Live? Yes/No | Customs, Other Apps | ...

Column H mapping:
    Customs / Customes -> Customs
    anything else (Other Apps, blank, ...) -> Internal IT

Examples:
    python scripts/import_applications_from_excel.py --excel portfolio.xlsx
    python scripts/import_applications_from_excel.py --excel portfolio.xlsx --apply
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import time
import urllib.error
import urllib.request
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from excel_import.excel_reader import (  # noqa: E402
    SheetDetection,
    _read_mapped_rows,
    detect_phase2_headers,
)
from excel_import.normalize import (  # noqa: E402
    cell_text,
    generate_application_code,
    is_blank,
    map_customs_or_internal_it,
    map_status,
    map_support_type,
    normalize_application_name,
    normalize_whitespace,
    parse_bool,
    split_technologies,
)

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover
    raise SystemExit(
        "openpyxl is required. Install it with:\n"
        "  pip install -r scripts/requirements-migration.txt"
    ) from exc


DEFAULT_API_URL = os.environ.get("API_BASE_URL", "http://127.0.0.1:8000/api/v1")
DEFAULT_AUTH_EMAIL = os.environ.get("AUTH_EMAIL", "admin@zatca.gov.sa")
DEFAULT_AUTH_PASSWORD = os.environ.get("AUTH_PASSWORD", "password")
DEFAULT_CRITICALITY = "Medium"

CUSTOMS_LOOKUP_NAMES = ["Customs", "Customes", "Custom"]
INTERNAL_IT_LOOKUP_NAMES = ["Internal IT", "Internal App", "Internal"]


@dataclass
class PreparedApplication:
    row: int
    name_en: str
    name_ar: str
    code: str
    classification: str
    technical_category: str | None
    description: str | None
    is_live: bool | None
    is_under_operation: bool | None
    status_name: str
    support_type_name: str
    vendor_name: str | None
    remarks: str | None
    technologies: list[str] = field(default_factory=list)
    skipped_reason: str | None = None

    def payload(self, ids: dict[str, Any]) -> dict[str, Any]:
        body: dict[str, Any] = {
            "department_id": ids["department_id"],
            "application_type_id": ids["application_type_id"],
            "name_ar": self.name_ar,
            "name_en": self.name_en,
            "code": self.code,
            "status_id": ids["status_id"],
            "criticality_id": ids["criticality_id"],
            "support_type_id": ids["support_type_id"],
        }
        if self.description:
            body["description"] = self.description
        if self.technical_category:
            body["technical_category"] = self.technical_category
        if self.remarks:
            body["remarks"] = self.remarks
        if ids.get("vendor_id"):
            body["vendor_id"] = ids["vendor_id"]
        if ids.get("technology_ids"):
            body["technologies"] = ids["technology_ids"]
        return body


class ApiError(RuntimeError):
    def __init__(self, status: int, message: str, errors: Any = None) -> None:
        super().__init__(message)
        self.status = status
        self.errors = errors


class ApplicationsApiClient:
    def __init__(self, base_url: str, timeout: int = 30) -> None:
        self.base_url = base_url.rstrip("/")
        self.timeout = timeout
        self.token: str | None = None

    def login(self, email: str, password: str) -> None:
        data = self._request(
            "POST",
            "/auth/login",
            {"email": email, "password": password},
            authenticated=False,
        )
        token = (data or {}).get("token")
        if not token:
            raise ApiError(0, "Login succeeded but no token was returned.")
        self.token = str(token)

    def lookup(self, path: str) -> list[dict[str, Any]]:
        data = self._request("GET", path)
        if isinstance(data, list):
            return data
        if isinstance(data, dict) and isinstance(data.get("items"), list):
            return data["items"]
        return []

    def paginated(self, path: str) -> list[dict[str, Any]]:
        items: list[dict[str, Any]] = []
        page = 1
        last_page = 1
        separator = "&" if "?" in path else "?"
        while page <= last_page:
            data = self._request("GET", f"{path}{separator}page={page}&per_page=100")
            if isinstance(data, list):
                return data
            chunk = (data or {}).get("items") or []
            pagination = (data or {}).get("pagination") or {}
            last_page = int(pagination.get("last_page") or 1)
            items.extend(chunk)
            page += 1
        return items

    def create_department(self, name_en: str, name_ar: str) -> dict[str, Any]:
        return self._request("POST", "/departments", {"name_en": name_en, "name_ar": name_ar}) or {}

    def create_application(self, payload: dict[str, Any]) -> dict[str, Any]:
        return self._request("POST", "/applications", payload) or {}

    def _request(
        self,
        method: str,
        path: str,
        body: dict[str, Any] | None = None,
        authenticated: bool = True,
    ) -> dict[str, Any] | list[Any] | None:
        url = f"{self.base_url}{path}"
        headers = {
            "Accept": "application/json",
            "Content-Type": "application/json",
        }
        if authenticated:
            if not self.token:
                raise ApiError(401, "Not authenticated. Login first.")
            headers["Authorization"] = f"Bearer {self.token}"

        encoded = None if body is None else json.dumps(body).encode("utf-8")
        request = urllib.request.Request(url, data=encoded, headers=headers, method=method)
        try:
            with urllib.request.urlopen(request, timeout=self.timeout) as response:
                raw = response.read().decode("utf-8")
        except urllib.error.HTTPError as exc:
            raw = exc.read().decode("utf-8", errors="replace")
            payload = _decode_json(raw)
            message = str(payload.get("message") or exc.reason or "API request failed")
            raise ApiError(exc.code, message, payload.get("errors")) from exc
        except urllib.error.URLError as exc:
            raise ApiError(0, f"Could not reach API at {url}: {exc.reason}") from exc

        payload = _decode_json(raw)
        if payload.get("success") is False:
            raise ApiError(
                0,
                str(payload.get("message") or "API request failed"),
                payload.get("errors"),
            )
        return payload.get("data")


def _decode_json(raw: str) -> dict[str, Any]:
    if not raw.strip():
        return {}
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        return {"message": raw}
    if isinstance(parsed, dict):
        return parsed
    return {"data": parsed}


def find_phase2_sheet(workbook, sheet_name: str | None = None):
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise RuntimeError(
                f"Sheet '{sheet_name}' was not found. Available: {workbook.sheetnames}"
            )
        return workbook[sheet_name]

    for name in workbook.sheetnames:
        if name.strip().casefold() == "phase2":
            return workbook[name]

    for name in workbook.sheetnames:
        sheet = workbook[name]
        try:
            detect_phase2_headers(sheet)
            return sheet
        except RuntimeError:
            continue

    raise RuntimeError("Could not detect a PHASE2 sheet. Expected an 'App Names' header.")


def extract_applications(
    path: Path,
    sheet_name: str | None = None,
    header_row: int | None = None,
    limit: int | None = None,
) -> tuple[list[PreparedApplication], SheetDetection]:
    workbook = load_workbook(filename=path, data_only=True, read_only=True)
    try:
        sheet = find_phase2_sheet(workbook, sheet_name)
        detection = detect_phase2_headers(sheet)
        if header_row is not None:
            detection = SheetDetection(detection.sheet_name, header_row, detection.columns, detection.warnings)

        rows = _read_mapped_rows(sheet, detection)
        applications: list[PreparedApplication] = []
        seen_names: set[str] = set()
        taken_codes: set[str] = set()

        for record in rows:
            name = normalize_whitespace(cell_text(record.get("app_name")))
            if name == "":
                applications.append(
                    PreparedApplication(
                        row=int(record.get("_row") or 0),
                        name_en="",
                        name_ar="",
                        code="",
                        classification="",
                        technical_category=None,
                        description=None,
                        is_live=None,
                        is_under_operation=None,
                        status_name="",
                        support_type_name="",
                        vendor_name=None,
                        remarks=None,
                        skipped_reason="missing application name",
                    )
                )
                continue

            name_key = normalize_application_name(name)
            classification = map_customs_or_internal_it(record.get("application_type"))
            technical_category = normalize_whitespace(cell_text(record.get("technical_category"))) or None
            description = normalize_whitespace(cell_text(record.get("description"))) or None
            vendor_name = normalize_whitespace(cell_text(record.get("vendor_name"))) or None
            remarks = normalize_whitespace(cell_text(record.get("remarks"))) or None
            live_raw = cell_text(record.get("live"))
            under_raw = cell_text(record.get("under_operation"))
            is_live = parse_bool(live_raw)
            is_under_operation = parse_bool(under_raw)
            if under_raw and is_under_operation is None:
                extra = normalize_whitespace(under_raw)
                remarks = extra if not remarks else f"{remarks}; Under Operation: {extra}"
                is_under_operation = False

            skipped_reason: str | None = None
            if name_key in seen_names:
                skipped_reason = "duplicate application name in sheet"

            seen_names.add(name_key)
            code = generate_application_code(name, taken_codes)
            taken_codes.add(code)

            applications.append(
                PreparedApplication(
                    row=int(record.get("_row") or 0),
                    name_en=name,
                    name_ar=name,
                    code=code,
                    classification=classification,
                    technical_category=technical_category,
                    description=description,
                    is_live=is_live,
                    is_under_operation=is_under_operation,
                    status_name=map_status(is_live),
                    support_type_name=map_support_type(is_under_operation),
                    vendor_name=vendor_name,
                    remarks=remarks,
                    technologies=split_technologies(cell_text(record.get("stack"))),
                    skipped_reason=skipped_reason,
                )
            )
            if limit is not None and len(applications) >= limit:
                break

        return applications, detection
    finally:
        workbook.close()


def match_lookup(items: list[dict[str, Any]], names: list[str]) -> dict[str, Any] | None:
    wanted = {normalize_whitespace(name).casefold() for name in names if name}
    for item in items:
        for key in ("name_en", "name_ar", "code", "name"):
            value = item.get(key)
            if value and normalize_whitespace(str(value)).casefold() in wanted:
                return item
    return None


def lookup_id(items: list[dict[str, Any]], names: list[str], label: str) -> int:
    match = match_lookup(items, names)
    if match is None or match.get("id") is None:
        raise ApiError(0, f"Could not resolve {label} from {names}. Seed lookups first.")
    return int(match["id"])


def classification_names(classification: str) -> list[str]:
    if classification == "Customs":
        return CUSTOMS_LOOKUP_NAMES
    return [classification, *INTERNAL_IT_LOOKUP_NAMES]


def format_errors(errors: Any) -> str:
    if not errors:
        return ""
    if isinstance(errors, dict):
        parts: list[str] = []
        for field, messages in errors.items():
            if isinstance(messages, list):
                parts.append(f"{field}: {'; '.join(str(item) for item in messages)}")
            else:
                parts.append(f"{field}: {messages}")
        return " | ".join(parts)
    return str(errors)


def print_preview(applications: list[PreparedApplication], detection: SheetDetection) -> None:
    ready = [item for item in applications if item.skipped_reason is None]
    skipped = [item for item in applications if item.skipped_reason is not None]
    print(
        f"Sheet '{detection.sheet_name}' header row {detection.header_row}: "
        f"{len(ready)} ready, {len(skipped)} skipped."
    )
    print()
    print(
        f"{'Row':<6} {'Code':<28} {'Classification':<14} {'Status':<12} "
        f"{'Support':<16} {'Name'}"
    )
    print("-" * 120)
    for item in applications:
        status = "OK" if item.skipped_reason is None else f"SKIP ({item.skipped_reason})"
        print(
            f"{item.row:<6} {item.code:<28} {item.classification:<14} {item.status_name:<12} "
            f"{item.support_type_name:<16} {item.name_en}  {status}"
        )


def resolve_department_id(client: ApplicationsApiClient, classification: str) -> int:
    departments = client.lookup("/lookups/departments")
    names = classification_names(classification)
    match = match_lookup(departments, names)
    if match is not None:
        return int(match["id"])

    created = client.create_department(classification, classification)
    department_id = created.get("id")
    if department_id is None:
        raise ApiError(0, f"Created department '{classification}' but no id was returned.")
    print(f"Created department '{classification}' #{department_id}")
    return int(department_id)


def resolve_application_ids(
    application: PreparedApplication,
    client: ApplicationsApiClient,
    cache: dict[str, Any],
) -> dict[str, Any]:
    classification = application.classification
    department_key = f"department:{classification}"
    type_key = f"type:{classification}"

    if department_key not in cache:
        cache[department_key] = resolve_department_id(client, classification)
    if type_key not in cache:
        types = cache.setdefault("application_types", client.lookup("/lookups/application-types"))
        names = classification_names(classification)
        match = match_lookup(types, names)
        if match is None:
            raise ApiError(
                0,
                f"Application type '{classification}' was not found. "
                "Create it in the UI or seed application types, then re-run.",
            )
        resolved = str(match.get("name_en") or classification)
        if normalize_whitespace(resolved).casefold() != classification.casefold():
            print(
                f"Note: '{classification}' is not an application type; using '{resolved}'."
            )
        cache[type_key] = int(match["id"])

    vendor_id = None
    if application.vendor_name:
        vendors = cache.setdefault("vendors", client.paginated("/vendors"))
        vendor = match_lookup(vendors, [application.vendor_name])
        if vendor is not None:
            vendor_id = int(vendor["id"])

    technology_ids: list[int] = []
    if application.technologies:
        technologies = cache.setdefault("technologies", client.lookup("/lookups/technologies"))
        for tech_name in application.technologies:
            tech = match_lookup(technologies, [tech_name])
            if tech is not None:
                technology_ids.append(int(tech["id"]))

    ids: dict[str, Any] = {
        "department_id": cache[department_key],
        "application_type_id": cache[type_key],
        "status_id": cache["status_ids"][application.status_name],
        "criticality_id": cache["criticality_id"],
        "support_type_id": cache["support_ids"][application.support_type_name],
    }
    if vendor_id is not None:
        ids["vendor_id"] = vendor_id
    if technology_ids:
        ids["technology_ids"] = technology_ids
    return ids


def import_applications(
    applications: list[PreparedApplication],
    client: ApplicationsApiClient,
    delay: float,
    skip_existing: bool,
) -> tuple[int, int, int]:
    created = skipped = failed = 0
    cache: dict[str, Any] = {
        "status_ids": {
            "Active": lookup_id(client.lookup("/lookups/application-statuses"), ["Active"], "status Active"),
            "Maintenance": lookup_id(
                client.lookup("/lookups/application-statuses"),
                ["Maintenance"],
                "status Maintenance",
            ),
        },
        "support_ids": {
            "Business Hours": lookup_id(
                client.lookup("/lookups/support-types"),
                ["Business Hours"],
                "support type Business Hours",
            ),
            "Best Effort": lookup_id(
                client.lookup("/lookups/support-types"),
                ["Best Effort"],
                "support type Best Effort",
            ),
        },
        "criticality_id": lookup_id(
            client.lookup("/lookups/criticalities"),
            [DEFAULT_CRITICALITY],
            f"criticality {DEFAULT_CRITICALITY}",
        ),
    }

    existing_names: set[str] = set()
    existing_codes: set[str] = set()
    if skip_existing:
        print("Loading existing applications from API...")
        for item in client.paginated("/applications"):
            name = normalize_application_name(str(item.get("name_en") or ""))
            code = str(item.get("code") or "").casefold()
            if name:
                existing_names.add(name)
            if code:
                existing_codes.add(code)
        print(f"Found {len(existing_names)} existing applications.")

    ready = [item for item in applications if item.skipped_reason is None]
    for index, application in enumerate(ready, start=1):
        name_key = normalize_application_name(application.name_en)
        if skip_existing and (name_key in existing_names or application.code.casefold() in existing_codes):
            print(f"[{index}/{len(ready)}] skip existing {application.name_en}")
            skipped += 1
            continue

        try:
            ids = resolve_application_ids(application, client, cache)
            created_app = client.create_application(application.payload(ids))
            app_id = created_app.get("id", "?")
            print(
                f"[{index}/{len(ready)}] created #{app_id} {application.code} "
                f"{application.name_en} [{application.classification}]"
            )
            created += 1
            existing_names.add(name_key)
            existing_codes.add(application.code.casefold())
        except ApiError as exc:
            failed += 1
            details = format_errors(exc.errors)
            suffix = f" ({details})" if details else ""
            print(
                f"[{index}/{len(ready)}] FAILED {application.name_en}: "
                f"HTTP {exc.status} {exc}{suffix}"
            )
        if delay > 0:
            time.sleep(delay)

    sheet_skipped = sum(1 for item in applications if item.skipped_reason is not None)
    return created, skipped + sheet_skipped, failed


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Extract applications from the PHASE2 Excel sheet and insert them through "
            "POST /api/v1/applications."
        )
    )
    parser.add_argument("--excel", required=True, type=Path, help="Path to the Excel workbook.")
    parser.add_argument("--sheet", default=None, help="Sheet name. Defaults to PHASE2 when present.")
    parser.add_argument("--header-row", type=int, default=None, help="1-based header row override.")
    parser.add_argument("--api-url", default=DEFAULT_API_URL, help=f"API base URL. Default: {DEFAULT_API_URL}")
    parser.add_argument("--auth-email", default=DEFAULT_AUTH_EMAIL, help="Account with applications.create permission.")
    parser.add_argument("--auth-password", default=DEFAULT_AUTH_PASSWORD, help="Login password for --auth-email.")
    parser.add_argument("--limit", type=int, default=None, help="Only process the first N data rows.")
    parser.add_argument("--delay", type=float, default=0.15, help="Seconds to wait between create calls.")
    parser.add_argument(
        "--apply",
        action="store_true",
        help="Create applications through the API. Without this flag the script only prints a dry run.",
    )
    parser.add_argument(
        "--no-skip-existing",
        action="store_true",
        help="Do not pre-load existing applications; still fails on unique-code validation errors.",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    excel_path = args.excel.expanduser().resolve()
    if not excel_path.is_file():
        print(f"Excel file not found: {excel_path}", file=sys.stderr)
        return 1

    applications, detection = extract_applications(
        excel_path,
        sheet_name=args.sheet,
        header_row=args.header_row,
        limit=args.limit,
    )
    print_preview(applications, detection)

    ready = [item for item in applications if item.skipped_reason is None]
    if not args.apply:
        print()
        print(
            f"Dry run only. Re-run with --apply to create {len(ready)} applications "
            f"via {args.api_url}/applications"
        )
        return 0

    if not ready:
        print("Nothing to import.")
        return 0

    client = ApplicationsApiClient(args.api_url)
    print()
    print(f"Logging in as {args.auth_email}...")
    try:
        client.login(args.auth_email, args.auth_password)
    except ApiError as exc:
        print(f"Login failed: {exc}", file=sys.stderr)
        return 1

    created, skipped, failed = import_applications(
        applications,
        client,
        delay=max(0.0, args.delay),
        skip_existing=not args.no_skip_existing,
    )
    print()
    print(f"Done. created={created} skipped={skipped} failed={failed}")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
