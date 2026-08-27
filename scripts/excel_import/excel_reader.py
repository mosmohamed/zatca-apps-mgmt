from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .normalize import cell_text, header_key, is_blank

USERS_REQUIRED_HEADERS = {
    "id",
    "display name",
    "email",
    "applications",
    "phone number",
}

USERS_HEADER_ALIASES = {
    "id": "id",
    "#": "id",
    "display name": "display_name",
    "name": "display_name",
    "employee name": "display_name",
    "email": "email",
    "e-mail": "email",
    "applications": "applications",
    "application": "applications",
    "apps": "applications",
    "phone number": "phone",
    "phone": "phone",
    "mobile": "phone",
}

PHASE2_HEADER_ALIASES = {
    "#": "excel_number",
    "no": "excel_number",
    "app names": "app_name",
    "app name": "app_name",
    "application name": "app_name",
    "application names": "app_name",
    "name": "app_name",
    "technical category": "technical_category",
    "category": "technical_category",
    "app description": "description",
    "description": "description",
    "under operation": "under_operation",
    "is it under operation": "under_operation",
    "under operation?": "under_operation",
    "live": "live",
    "live? yes/no": "live",
    "live yes/no": "live",
    "customs, other apps": "application_type",
    "customs other apps": "application_type",
    "customs, other app": "application_type",
    "application type": "application_type",
    "support name": "support_name",
    "who support": "support_name",
    "who support?": "support_name",
    "support email": "support_email",
    "email": "support_email",
    "vendor name": "vendor_name",
    "vendor": "vendor_name",
    "zatca management owner": "management_owner",
    "management owner": "management_owner",
    "management owner email": "management_owner_email",
    "zatca app lead": "app_lead",
    "app lead": "app_lead",
    "app lead email": "app_lead_email",
    "stack and technologies": "stack",
    "stack": "stack",
    "technologies": "stack",
    "remarks": "remarks",
    "remark": "remarks",
}


@dataclass
class SheetDetection:
    sheet_name: str
    header_row: int
    columns: dict[str, int]
    warnings: list[str] = field(default_factory=list)


@dataclass
class ExcelPayload:
    users_rows: list[dict[str, Any]]
    phase2_rows: list[dict[str, Any]]
    users_detection: SheetDetection
    phase2_detection: SheetDetection
    sheet_names: list[str]


def load_workbook_payload(path: Path) -> ExcelPayload:
    workbook = load_workbook(filename=path, data_only=True, read_only=True)
    try:
        sheet_names = list(workbook.sheetnames)
        phase2_sheet = _find_phase2_sheet(workbook)
        users_sheet = _find_users_sheet(workbook, exclude={phase2_sheet.title})

        phase2_detection = detect_phase2_headers(phase2_sheet)
        users_detection = detect_users_headers(users_sheet)

        return ExcelPayload(
            users_rows=_read_mapped_rows(users_sheet, users_detection),
            phase2_rows=_read_mapped_rows(phase2_sheet, phase2_detection),
            users_detection=users_detection,
            phase2_detection=phase2_detection,
            sheet_names=sheet_names,
        )
    finally:
        workbook.close()


def _find_phase2_sheet(workbook) -> Worksheet:
    for name in workbook.sheetnames:
        if name.strip().casefold() == "phase2":
            return workbook[name]
    raise RuntimeError("Workbook does not contain a sheet named PHASE2.")


def _find_users_sheet(workbook, exclude: set[str]) -> Worksheet:
    candidates: list[tuple[int, Worksheet, SheetDetection]] = []
    for name in workbook.sheetnames:
        if name in exclude:
            continue
        sheet = workbook[name]
        detection = detect_users_headers(sheet, required_ratio=0.8)
        score = len(set(detection.columns) & {"display_name", "email", "applications", "phone"})
        if score >= 3:
            candidates.append((score, sheet, detection))

    if not candidates:
        raise RuntimeError(
            "Could not detect a Users sheet. Expected headers: "
            "id, Display Name, Email, Applications, Phone Number."
        )

    candidates.sort(key=lambda item: (-item[0], item[1].title))
    return candidates[0][1]


def detect_users_headers(sheet: Worksheet, required_ratio: float = 1.0) -> SheetDetection:
    best: SheetDetection | None = None
    best_score = -1

    max_row = min(sheet.max_row or 1, 15)
    for row_idx in range(1, max_row + 1):
        mapping: dict[str, int] = {}
        warnings: list[str] = []
        for col_idx in range(1, min((sheet.max_column or 1), 20) + 1):
            key = header_key(sheet.cell(row_idx, col_idx).value)
            if key == "":
                continue
            field_name = USERS_HEADER_ALIASES.get(key)
            if field_name is None:
                continue
            if field_name in mapping:
                warnings.append(
                    f"Duplicate users header '{key}' on {sheet.title} row {row_idx}."
                )
                continue
            mapping[field_name] = col_idx

        score = len(mapping)
        required_hits = len({"display_name", "email", "applications"} & set(mapping))
        if required_hits < 3:
            continue
        if score > best_score:
            best_score = score
            best = SheetDetection(sheet.title, row_idx, mapping, warnings)

    if best is None:
        return SheetDetection(sheet.title, 2, {}, ["Users headers were not detected."])

    present = set(best.columns)
    missing = {"display_name", "email", "applications"} - present
    if missing and required_ratio >= 1:
        best.warnings.append(f"Users sheet is missing headers: {sorted(missing)}")
    return best


def detect_phase2_headers(sheet: Worksheet) -> SheetDetection:
    best: SheetDetection | None = None
    best_score = -1
    max_row = min(sheet.max_row or 1, 8)

    for row_idx in range(1, max_row + 1):
        mapping: dict[str, int] = {}
        warnings: list[str] = []
        seen_email_slots = 0
        for col_idx in range(1, min((sheet.max_column or 1), 25) + 1):
            raw = cell_text(sheet.cell(row_idx, col_idx).value)
            key = header_key(raw)
            if key == "":
                continue

            field_name = PHASE2_HEADER_ALIASES.get(key)
            if field_name is None and key.startswith("email"):
                seen_email_slots += 1
                field_name = {
                    1: "support_email",
                    2: "management_owner_email",
                    3: "app_lead_email",
                }.get(seen_email_slots)
            elif key == "email":
                seen_email_slots += 1
                field_name = {
                    1: "support_email",
                    2: "management_owner_email",
                    3: "app_lead_email",
                }.get(seen_email_slots, "support_email")

            if field_name is None:
                continue
            if field_name in mapping:
                warnings.append(
                    f"Duplicate PHASE2 header '{raw}' on {sheet.title} row {row_idx}."
                )
                continue
            mapping[field_name] = col_idx

        score = len(mapping)
        if "app_name" in mapping and score > best_score:
            best_score = score
            best = SheetDetection(sheet.title, row_idx, mapping, warnings)

    if best is None:
        raise RuntimeError("PHASE2 sheet is missing an 'App Names' header.")

    if "app_name" not in best.columns:
        raise RuntimeError("PHASE2 sheet is missing the App Names column.")

    return best


def _read_mapped_rows(sheet: Worksheet, detection: SheetDetection) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    start = detection.header_row + 1
    max_row = sheet.max_row or start

    for row_idx in range(start, max_row + 1):
        record: dict[str, Any] = {"_row": row_idx, "_sheet": detection.sheet_name}
        empty = True
        for field_name, col_idx in detection.columns.items():
            value = sheet.cell(row_idx, col_idx).value
            record[field_name] = value
            if not is_blank(value):
                empty = False
        if empty:
            continue
        rows.append(record)
    return rows
