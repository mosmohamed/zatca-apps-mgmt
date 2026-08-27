#!/usr/bin/env python3
"""Generate a self-contained Laravel seeder from Users + PHASE2 Excel files.

Excel processing happens here. The generated PHP seeder contains the imported
data as arrays and never reads Excel, CSV, JSON, or any other external file.

Usage:
    python excel_to_laravel_seeder.py users.xlsx applications.xlsx
    python excel_to_laravel_seeder.py users.xlsx applications.xlsx --output ImportedPortfolioSeeder.php
    python excel_to_laravel_seeder.py users.xlsx applications.xlsx --output ImportedPortfolioSeeder.php --report migration_report.txt

Dependencies:
    pip install openpyxl mysql-connector-python

Applications are created only from the PHASE2 sheet. Users-sheet app names that
do not match a PHASE2 application still create the user, but skip that assignment.

MySQL is inspected read-only (SELECT / SHOW only) so the seeder can look up
related records by stable names and emails. No IDs are hardcoded.
"""

from __future__ import annotations

import argparse
import json
import math
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
    from openpyxl.worksheet.worksheet import Worksheet
except ImportError as exc:  # pragma: no cover
    raise SystemExit(
        "Missing dependency 'openpyxl'. Install with: pip install openpyxl"
    ) from exc


EMAIL_DOMAIN = "zatca.gov.sa"
EMAIL_RE = re.compile(r"^[A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,}$", re.IGNORECASE)
WHITESPACE_RE = re.compile(r"\s+", re.UNICODE)
UNSUPPORTED_LOCAL_RE = re.compile(r"[^a-z0-9._+\-]+")
SUPPORT_NAME_SPLIT_RE = re.compile(r"\s+[-–—]\s+")
EMAIL_SPLIT_RE = re.compile(r"\s*(?:[-–—]|[,;]|\r?\n)+\s*")
APP_SPLIT_RE = re.compile(r"\s*[,;]+\s*|\r?\n+")
TECH_SPLIT_RE = re.compile(r"\s*[,;/|]+\s*|\r?\n+")

PLACEHOLDER_REVIEW = "TODO - Needs Review"
PLACEHOLDER_VENDOR = "Unknown Vendor"
PLACEHOLDER_TECHNICAL_CATEGORY = "Unknown Technical Category"
PLACEHOLDER_PHONE = "0000000000"
PLACEHOLDER_APPLICATION_TYPE = "Internal App"

DEFAULT_DB = {
    "host": "127.0.0.1",
    "port": 3306,
    "database": "it_portfolio_system",
    "user": "root",
    "password": "",
}

USERS_EXPECTED_HEADERS = ["id", "Display Name", "Email", "Applications", "Phone Number"]
USERS_HEADER_ALIASES = {
    "id": "id",
    "#": "id",
    "display name": "display_name",
    "name": "display_name",
    "employee name": "display_name",
    "email": "email",
    "e-mail": "email",
    "applications": "applications",
    "application": "applications",
    "apps": "applications",
    "phone number": "phone",
    "phone": "phone",
    "mobile": "phone",
}

# PHASE2 B1:Q1 in order.
PHASE2_EXPECTED_HEADERS = [
    "#",
    "App Names",
    "Technical Category",
    "App Description",
    "under Operation",
    "Live",
    "Customs, Other Apps",
    "Support Name",
    "Email",
    "Vendor Name",
    "Zatca Management Owner",
    "Email",
    "Zatca App Lead",
    "Email",
    "Stack and Technologies",
    "Remarks",
]
PHASE2_FIELD_ORDER = [
    "excel_number",
    "app_name",
    "technical_category",
    "description",
    "under_operation",
    "live",
    "application_type",
    "support_name",
    "support_email",
    "vendor_name",
    "management_owner",
    "management_owner_email",
    "app_lead",
    "app_lead_email",
    "stack",
    "remarks",
]
PHASE2_HEADER_START_COL = 2  # B
PHASE2_HEADER_END_COL = 17  # Q

APPLICATION_TYPE_ALIASES = {
    "other apps": "Internal App",
    "other app": "Internal App",
    "others": "Internal App",
    "other": "Internal App",
    "internal": "Internal App",
    "internal app": "Internal App",
    "internal apps": "Internal App",
    "customs": "Customs",
    "customes": "Customs",
    "custom": "Customs",
}

APP_ROLE_RANK = {
    "ZATCA Management": 40,
    "Application Lead": 30,
    "Support": 20,
    "Viewer": 10,
}

STATIC_USERS = [
    {
        "first_name": "Super",
        "last_name": "Admin",
        "email": "super_admin@zatca.gov.sa",
        "phone": "5678910110",
        "role": "super_admin",
        "display_name": "super admin",
    },
    {
        "first_name": "Infra",
        "last_name": "Admin",
        "email": "infra_admin@zatca.gov.sa",
        "phone": "5678910110",
        "role": "infra_admin",
        "display_name": "infra admin",
    },
    {
        "first_name": "SD",
        "last_name": "Admin",
        "email": "sd_admin@zatca.gov.sa",
        "phone": "5678910110",
        "role": "sd_admin",
        "display_name": "SD admin",
    },
    {
        "first_name": "Viewer",
        "last_name": "User",
        "email": "viewer@zatca.gov.sa",
        "phone": "5678910110",
        "role": "viewer",
        "display_name": "viewer user",
    },
]

FALLBACK_APPLICATION_TYPES = [
    "API",
    "Web Portal",
    "Mobile App",
    "Microservice",
    "Desktop Application",
    "SaaS Platform",
    "Customs",
    "Internal App",
    "Application",
    "Tool",
    "Service",
    "AI Model",
]
FALLBACK_DEPARTMENTS = ["Customes", "General IT", "taxation and Zakat Department"]
FALLBACK_APP_ROLES = [
    "Admin",
    "Developer",
    "QA",
    "Support",
    "ZATCA Management",
    "Viewer",
    "Application Lead",
]
FALLBACK_SPATIE_ROLES = ["super_admin", "infra_admin", "sd_admin", "viewer", "employee"]

TRUTHY = {"yes", "y", "true", "live", "1", "on", "active"}
FALSY = {"no", "n", "false", "not live", "notlive", "0", "off", "inactive"}


# ---------------------------------------------------------------------------
# Normalization
# ---------------------------------------------------------------------------


def is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    text = str(value).strip()
    return text == "" or text.lower() in {"nan", "none", "null", "-", "n/a", "na"}


def cell_text(value: Any) -> str:
    if is_blank(value):
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).strip()
    if isinstance(value, int):
        return str(value)
    return str(value).replace("\xa0", " ").strip()


def normalize_whitespace(value: str) -> str:
    return WHITESPACE_RE.sub(" ", value.replace("\xa0", " ")).strip()


def normalize_name(value: str) -> str:
    return normalize_whitespace(value).casefold()


def normalize_email(value: str) -> str:
    return normalize_whitespace(value).lower().replace(" ", "")


def normalize_application_name(value: str) -> str:
    return normalize_whitespace(value).casefold()


def header_key(value: Any) -> str:
    text = normalize_whitespace(cell_text(value)).lower()
    text = re.sub(r"[?]+$", "", text)
    text = text.replace("&", "and")
    return WHITESPACE_RE.sub(" ", text)


def headers_equivalent(actual: str, expected: str) -> bool:
    left = re.sub(r"[^\w]+", " ", header_key(actual)).strip()
    right = re.sub(r"[^\w]+", " ", header_key(expected)).strip()
    return left == right


def split_display_name(display_name: str) -> tuple[str, str]:
    parts = [part for part in normalize_whitespace(display_name).split(" ") if part]
    if not parts:
        return "", ""
    if len(parts) == 1:
        return parts[0], parts[0]
    return parts[0], " ".join(parts[1:])


def slug_local_part(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value)
    ascii_only = decomposed.encode("ascii", "ignore").decode("ascii")
    local = ascii_only.lower().replace(" ", "")
    local = UNSUPPORTED_LOCAL_RE.sub("", local)
    return local.strip("._+-")


def generate_email(first_name: str, last_name: str, domain: str = EMAIL_DOMAIN) -> str:
    first = slug_local_part(first_name)
    last = slug_local_part(last_name)
    if not first and not last:
        return f"user@{domain}"
    if not first:
        local = last
    elif not last:
        local = f"{first[0]}{first}"
    else:
        local = f"{first[0]}{last}"
    if not local:
        local = "user"
    return f"{local}@{domain}"


def unique_generated_email(base_email: str, taken: set[str]) -> tuple[str, bool]:
    candidate = normalize_email(base_email)
    if candidate not in taken:
        return candidate, False
    local, _, domain = candidate.partition("@")
    suffix = 2
    while True:
        next_candidate = f"{local}{suffix}@{domain}"
        if next_candidate not in taken:
            return next_candidate, True
        suffix += 1


def is_valid_email(value: str) -> bool:
    return bool(value) and EMAIL_RE.match(value) is not None


def parse_phone(value: Any) -> str | None:
    if is_blank(value):
        return None
    if isinstance(value, float):
        if math.isnan(value):
            return None
        if value.is_integer():
            return str(int(value))
        text = format(value, "f").rstrip("0").rstrip(".")
        return text or None
    if isinstance(value, int):
        return str(value)

    text = cell_text(value)
    if re.fullmatch(r"-?\d+\.0", text):
        text = text[:-2]
    text = text.replace(" ", "")
    return text or None


def split_support_names(value: str) -> list[str]:
    if is_blank(value):
        return []
    parts = SUPPORT_NAME_SPLIT_RE.split(normalize_whitespace(cell_text(value)))
    names: list[str] = []
    seen: set[str] = set()
    for part in parts:
        name = normalize_whitespace(part)
        if name == "" or name in {"-", ",", ";"}:
            continue
        key = normalize_name(name)
        if key in seen:
            continue
        seen.add(key)
        names.append(name)
    return names


def split_emails(value: str) -> list[str]:
    if is_blank(value):
        return []
    emails: list[str] = []
    seen: set[str] = set()
    for part in EMAIL_SPLIT_RE.split(cell_text(value)):
        email = normalize_email(part)
        if not is_valid_email(email) or email in seen:
            continue
        seen.add(email)
        emails.append(email)
    return emails


def split_applications(value: str) -> list[str]:
    if is_blank(value):
        return []
    names: list[str] = []
    seen: set[str] = set()
    for part in APP_SPLIT_RE.split(cell_text(value)):
        name = normalize_whitespace(part)
        if name == "":
            continue
        key = normalize_application_name(name)
        if key in seen:
            continue
        seen.add(key)
        names.append(name)
    return names


def split_technologies(value: str) -> list[str]:
    if is_blank(value):
        return []
    names: list[str] = []
    seen: set[str] = set()
    for part in TECH_SPLIT_RE.split(cell_text(value)):
        name = normalize_whitespace(part)
        if name == "" or name.lower() in {"n/a", "na", "none", "nil"}:
            continue
        key = name.casefold()
        if key in seen:
            continue
        seen.add(key)
        names.append(name)
    return names


def parse_bool(value: Any) -> bool | None:
    if is_blank(value):
        return None
    if isinstance(value, bool):
        return value
    text = normalize_whitespace(cell_text(value)).casefold().replace("_", " ")
    if text in TRUTHY:
        return True
    if text in FALSY:
        return False
    return None


def preferred_display_name(current: str, incoming: str) -> str:
    if not current:
        return incoming
    if not incoming:
        return current
    if incoming != incoming.upper() and current == current.upper():
        return incoming
    if " " in incoming and " " not in current:
        return incoming
    return current


def php_class_name(path: Path) -> str:
    name = path.stem
    if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", name):
        return name
    return "ImportedPortfolioSeeder"


# ---------------------------------------------------------------------------
# Schema inspection (read-only)
# ---------------------------------------------------------------------------


@dataclass
class SchemaSnapshot:
    connected: bool
    database: str | None
    tables: list[str] = field(default_factory=list)
    columns: dict[str, list[str]] = field(default_factory=dict)
    users: list[dict[str, Any]] = field(default_factory=list)
    applications: list[dict[str, Any]] = field(default_factory=list)
    vendors: list[str] = field(default_factory=list)
    departments: list[str] = field(default_factory=list)
    application_types: list[str] = field(default_factory=list)
    technologies: list[str] = field(default_factory=list)
    app_roles: list[str] = field(default_factory=list)
    roles: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def preferred_application_name(self, name: str) -> str:
        key = normalize_application_name(name)
        for existing in self.applications:
            existing_name = str(existing.get("name_en") or "")
            if normalize_application_name(existing_name) == key:
                return existing_name
        return name

    def preferred_vendor_name(self, name: str) -> str:
        key = name.casefold()
        for existing in self.vendors:
            if existing.casefold() == key:
                return existing
        return name

    def preferred_application_type(self, name: str) -> str:
        key = name.casefold()
        for existing in self.application_types:
            if existing.casefold() == key:
                return existing
        return name

    def preferred_department_name(self, name: str) -> str:
        key = name.casefold()
        for existing in self.departments:
            if existing.casefold() == key:
                return existing
        if key in {"customs", "custom", "customes"}:
            for existing in self.departments:
                if "custom" in existing.casefold():
                    return existing
        if key in {"internal", "internal app", "other apps", "taxation and zakat department"}:
            for existing in self.departments:
                lowered = existing.casefold()
                if "zakat" in lowered or "tax" in lowered:
                    return existing
        return name

    def user_email_by_name(self, full_name: str) -> str | None:
        key = normalize_name(full_name)
        for user in self.users:
            combined = normalize_whitespace(
                f"{user.get('first_name') or ''} {user.get('last_name') or ''}"
            )
            if normalize_name(combined) == key:
                email = user.get("email")
                return str(email) if email else None
        return None


def inspect_schema(
    host: str,
    port: int,
    database: str,
    user: str,
    password: str,
    skip_db: bool,
) -> SchemaSnapshot:
    snapshot = SchemaSnapshot(
        connected=False,
        database=database,
        application_types=list(FALLBACK_APPLICATION_TYPES),
        departments=list(FALLBACK_DEPARTMENTS),
        app_roles=list(FALLBACK_APP_ROLES),
        roles=list(FALLBACK_SPATIE_ROLES),
    )
    if skip_db:
        snapshot.warnings.append("Database inspection skipped (--skip-db). Using schema fallbacks.")
        return snapshot

    try:
        import mysql.connector  # type: ignore
    except ImportError:
        snapshot.warnings.append(
            "mysql-connector-python is not installed. "
            "Install it with: pip install mysql-connector-python. Using schema fallbacks."
        )
        return snapshot

    try:
        connection = mysql.connector.connect(
            host=host,
            port=port,
            user=user,
            password=password,
            database=database,
            connection_timeout=5,
        )
    except Exception as exc:  # noqa: BLE001
        snapshot.warnings.append(f"MySQL inspection failed: {exc}. Using schema fallbacks.")
        return snapshot

    try:
        cursor = connection.cursor(dictionary=True)
        cursor.execute("SHOW TABLES")
        snapshot.tables = [str(next(iter(row.values()))) for row in cursor.fetchall()]
        snapshot.connected = True

        for table in snapshot.tables:
            cursor.execute(f"SHOW COLUMNS FROM `{table}`")
            snapshot.columns[table] = [str(row["Field"]) for row in cursor.fetchall()]

        if "users" in snapshot.tables:
            cursor.execute(
                "SELECT id, first_name, last_name, email, phone FROM users WHERE deleted_at IS NULL"
            )
            snapshot.users = list(cursor.fetchall())
        if "applications" in snapshot.tables:
            cursor.execute(
                "SELECT id, name_en, name_ar, code FROM applications WHERE deleted_at IS NULL"
            )
            snapshot.applications = list(cursor.fetchall())
        if "vendors" in snapshot.tables:
            cursor.execute("SELECT name FROM vendors WHERE deleted_at IS NULL")
            snapshot.vendors = [str(row["name"]) for row in cursor.fetchall() if row.get("name")]
        if "departments" in snapshot.tables:
            cursor.execute("SELECT name_en FROM departments WHERE deleted_at IS NULL")
            snapshot.departments = [
                str(row["name_en"]) for row in cursor.fetchall() if row.get("name_en")
            ]
        if "application_types" in snapshot.tables:
            cursor.execute("SELECT name_en FROM application_types")
            snapshot.application_types = [
                str(row["name_en"]) for row in cursor.fetchall() if row.get("name_en")
            ]
        if "technologies" in snapshot.tables:
            cursor.execute("SELECT name FROM technologies WHERE deleted_at IS NULL")
            snapshot.technologies = [
                str(row["name"]) for row in cursor.fetchall() if row.get("name")
            ]
        if "app_roles" in snapshot.tables:
            cursor.execute("SELECT name FROM app_roles WHERE deleted_at IS NULL")
            snapshot.app_roles = [str(row["name"]) for row in cursor.fetchall() if row.get("name")]
        if "roles" in snapshot.tables:
            cursor.execute("SELECT name FROM roles")
            snapshot.roles = [str(row["name"]) for row in cursor.fetchall() if row.get("name")]
        cursor.close()
    finally:
        connection.close()

    return snapshot


# ---------------------------------------------------------------------------
# Excel readers
# ---------------------------------------------------------------------------


@dataclass
class SheetDetection:
    sheet_name: str
    header_row: int
    columns: dict[str, int]
    warnings: list[str] = field(default_factory=list)


class HeaderMismatchError(RuntimeError):
    pass


def _row_values(sheet: Worksheet, row_idx: int, start_col: int, end_col: int) -> list[str]:
    return [cell_text(sheet.cell(row_idx, col_idx).value) for col_idx in range(start_col, end_col + 1)]


def validate_phase2_headers(sheet: Worksheet) -> SheetDetection:
    actual = _row_values(sheet, 1, PHASE2_HEADER_START_COL, PHASE2_HEADER_END_COL)
    expected = PHASE2_EXPECTED_HEADERS
    mismatches: list[str] = []
    for index, (got, want) in enumerate(zip(actual, expected, strict=True)):
        col_letter = chr(ord("B") + index)
        if got == "" or not headers_equivalent(got, want):
            mismatches.append(f"  {col_letter}1: expected '{want}', found '{got or '(empty)'}'")

    if mismatches:
        raise HeaderMismatchError(
            "PHASE2 headers do not match the expected B1:Q1 structure:\n"
            + "\n".join(mismatches)
            + "\nExpected: "
            + " | ".join(expected)
        )

    columns = {
        field_name: PHASE2_HEADER_START_COL + index
        for index, field_name in enumerate(PHASE2_FIELD_ORDER)
    }
    return SheetDetection(sheet.title, 1, columns)


def detect_users_headers(sheet: Worksheet) -> SheetDetection | None:
    # Preferred location: B2:F2.
    preferred = _row_values(sheet, 2, 2, 6)
    if preferred and all(
        headers_equivalent(got, want) for got, want in zip(preferred, USERS_EXPECTED_HEADERS, strict=True)
    ):
        return SheetDetection(
            sheet.title,
            2,
            {
                "id": 2,
                "display_name": 3,
                "email": 4,
                "applications": 5,
                "phone": 6,
            },
        )

    max_row = min(sheet.max_row or 1, 20)
    max_col = min(sheet.max_column or 1, 20)
    best: SheetDetection | None = None
    best_score = -1
    for row_idx in range(1, max_row + 1):
        mapping: dict[str, int] = {}
        for col_idx in range(1, max_col + 1):
            key = header_key(sheet.cell(row_idx, col_idx).value)
            field_name = USERS_HEADER_ALIASES.get(key)
            if field_name and field_name not in mapping:
                mapping[field_name] = col_idx
        score = len({"display_name", "email", "applications", "phone"} & set(mapping))
        if score >= 3 and score > best_score:
            best_score = score
            best = SheetDetection(sheet.title, row_idx, mapping)
    return best


def load_users_rows(path: Path) -> tuple[list[dict[str, Any]], SheetDetection]:
    workbook = load_workbook(filename=path, data_only=True, read_only=False)
    try:
        detections: list[SheetDetection] = []
        for sheet_name in workbook.sheetnames:
            detection = detect_users_headers(workbook[sheet_name])
            if detection is not None:
                detections.append(detection)
        if not detections:
            raise HeaderMismatchError(
                f"Could not locate a Users sheet in '{path.name}'. "
                "Expected headers at B2:F2: id, Display Name, Email, Applications, Phone Number."
            )
        detections.sort(key=lambda item: (-len(item.columns), item.sheet_name))
        detection = detections[0]
        sheet = workbook[detection.sheet_name]
        rows = _read_mapped_rows(sheet, detection)
        return rows, detection
    finally:
        workbook.close()


def load_phase2_rows(path: Path) -> tuple[list[dict[str, Any]], SheetDetection, list[str]]:
    workbook = load_workbook(filename=path, data_only=True, read_only=False)
    try:
        sheet_names = list(workbook.sheetnames)
        phase2_name = next((name for name in sheet_names if name.strip().casefold() == "phase2"), None)
        if phase2_name is None:
            raise HeaderMismatchError(
                f"Applications workbook '{path.name}' does not contain a sheet named PHASE2. "
                f"Found sheets: {', '.join(sheet_names) or '(none)'}."
            )
        sheet = workbook[phase2_name]
        detection = validate_phase2_headers(sheet)
        return _read_mapped_rows(sheet, detection), detection, sheet_names
    finally:
        workbook.close()


def _read_mapped_rows(sheet: Worksheet, detection: SheetDetection) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    start = detection.header_row + 1
    max_row = sheet.max_row or start
    for row_idx in range(start, max_row + 1):
        record: dict[str, Any] = {"_row": row_idx, "_sheet": detection.sheet_name}
        empty = True
        for field_name, col_idx in detection.columns.items():
            value = sheet.cell(row_idx, col_idx).value
            record[field_name] = value
            if not is_blank(value):
                empty = False
        if empty:
            continue
        rows.append(record)
    return rows


# ---------------------------------------------------------------------------
# Canonical model + pipeline
# ---------------------------------------------------------------------------


@dataclass
class Placeholder:
    entity_type: str
    entity: str
    field: str
    value: str
    reason: str


@dataclass
class CanonicalUser:
    first_name: str
    last_name: str
    email: str
    phone: str | None = None
    display_name: str = ""
    email_source: str = "generated"
    spatie_role: str = "employee"
    sources: list[str] = field(default_factory=list)
    generated_email: bool = False
    static: bool = False


@dataclass
class CanonicalApplication:
    name: str
    application_type: str
    department: str
    technical_category: str | None = None
    description: str | None = None
    vendor: str | None = None
    support_type: str = "Business Hours"
    status: str = "Active"
    criticality: str = "Medium"
    technologies: list[str] = field(default_factory=list)
    remarks: str | None = None
    source: str = "phase2"
    management_owner_emails: list[str] = field(default_factory=list)
    application_lead_emails: list[str] = field(default_factory=list)
    placeholder_fields: list[str] = field(default_factory=list)


@dataclass
class CanonicalAssignment:
    user_email: str
    application_name: str
    app_role: str
    is_primary: bool
    source: str
    remarks: str | None = None


@dataclass
class ImportResult:
    users: list[CanonicalUser]
    applications: list[CanonicalApplication]
    assignments: list[CanonicalAssignment]
    placeholders: list[Placeholder]
    warnings: list[str]
    skipped_rows: list[dict[str, Any]]
    generated_emails: list[dict[str, str]]
    email_collisions: list[dict[str, str]]
    merged_users: list[str]
    merged_applications: list[str]
    unmatched_user_applications: list[dict[str, Any]]
    stats: dict[str, Any]


class ImportPipeline:
    def __init__(
        self,
        users_rows: list[dict[str, Any]],
        phase2_rows: list[dict[str, Any]],
        schema: SchemaSnapshot,
    ) -> None:
        self.users_rows = users_rows
        self.phase2_rows = phase2_rows
        self.schema = schema
        self.users: dict[str, CanonicalUser] = {}
        self.users_by_name: dict[str, str] = {}
        self.applications: dict[str, CanonicalApplication] = {}
        self.assignments: dict[tuple[str, str], CanonicalAssignment] = {}
        self.placeholders: list[Placeholder] = []
        self.warnings: list[str] = []
        self.skipped_rows: list[dict[str, Any]] = []
        self.generated_emails: list[dict[str, str]] = []
        self.email_collisions: list[dict[str, str]] = []
        self.merged_users: list[str] = []
        self.merged_applications: list[str] = []
        self.unmatched_user_applications: list[dict[str, Any]] = []
        self.taken_emails: set[str] = set()
        self.stats: dict[str, Any] = {
            "users_rows_processed": 0,
            "phase2_rows_processed": 0,
            "unique_users": 0,
            "users_merged": 0,
            "applications": 0,
            "assignments": 0,
            "generated_emails": 0,
            "email_collisions": 0,
            "placeholder_values": 0,
            "invalid_skipped_rows": 0,
            "unmatched_user_applications": 0,
        }

    def run(self) -> ImportResult:
        self._seed_existing_emails()
        self._seed_static_users()
        self._import_phase2()
        self._import_users_sheet()
        self.stats["unique_users"] = len(self.users)
        self.stats["users_merged"] = len(set(self.merged_users))
        self.stats["applications"] = len(self.applications)
        self.stats["assignments"] = len(self.assignments)
        self.stats["generated_emails"] = len(self.generated_emails)
        self.stats["email_collisions"] = len(self.email_collisions)
        self.stats["placeholder_values"] = len(self.placeholders)
        self.stats["invalid_skipped_rows"] = len(self.skipped_rows)
        self.stats["unmatched_user_applications"] = len(self.unmatched_user_applications)
        return ImportResult(
            users=sorted(self.users.values(), key=lambda item: item.email),
            applications=sorted(self.applications.values(), key=lambda item: item.name.casefold()),
            assignments=sorted(
                self.assignments.values(),
                key=lambda item: (item.application_name.casefold(), item.user_email, item.app_role),
            ),
            placeholders=self.placeholders,
            warnings=self.warnings,
            skipped_rows=self.skipped_rows,
            generated_emails=self.generated_emails,
            email_collisions=self.email_collisions,
            merged_users=sorted(set(self.merged_users)),
            merged_applications=sorted(set(self.merged_applications)),
            unmatched_user_applications=self.unmatched_user_applications,
            stats=self.stats,
        )

    def _seed_existing_emails(self) -> None:
        for user in self.schema.users:
            email = normalize_email(str(user.get("email") or ""))
            if email:
                self.taken_emails.add(email)

    def _seed_static_users(self) -> None:
        for row in STATIC_USERS:
            self._upsert_user(
                CanonicalUser(
                    first_name=str(row["first_name"]),
                    last_name=str(row["last_name"]),
                    email=str(row["email"]),
                    phone=str(row["phone"]),
                    display_name=str(row["display_name"]),
                    email_source="static",
                    spatie_role=str(row["role"]),
                    sources=["static"],
                    static=True,
                )
            )

    def _import_users_sheet(self) -> None:
        for row in self.users_rows:
            self.stats["users_rows_processed"] += 1
            display = normalize_whitespace(cell_text(row.get("display_name")))
            email_raw = normalize_email(cell_text(row.get("email")))
            if display == "" and email_raw == "":
                self._skip(row, "blank user")
                continue
            if display == "":
                self._skip(row, "missing Display Name")
                continue

            first, last = split_display_name(display)
            email = email_raw
            generated = False
            if email and not is_valid_email(email):
                self.warnings.append(
                    f"Users sheet row {row.get('_row')} has invalid email '{email}'; generating a fallback."
                )
                email = ""
            if not email:
                email, generated = self._make_generated_email(first, last, display)

            existing_before = self._find_user(display, email) is not None
            user = self._upsert_user(
                CanonicalUser(
                    first_name=first,
                    last_name=last,
                    email=email,
                    phone=parse_phone(row.get("phone")),
                    display_name=display,
                    email_source="users_sheet" if not generated else "generated",
                    spatie_role="employee",
                    sources=["users_sheet"],
                    generated_email=generated,
                )
            )
            if existing_before:
                self.merged_users.append(user.email)

            for app_name in split_applications(cell_text(row.get("applications"))):
                application = self._find_phase2_application(app_name)
                if application is None:
                    self.unmatched_user_applications.append(
                        {
                            "sheet": row.get("_sheet"),
                            "row": row.get("_row"),
                            "user": user.email,
                            "application": app_name,
                        }
                    )
                    self.warnings.append(
                        f"Users sheet row {row.get('_row')}: application '{app_name}' "
                        f"for {user.email} is not in PHASE2; creating the user without that assignment."
                    )
                    continue
                self._add_assignment(
                    CanonicalAssignment(
                        user_email=user.email,
                        application_name=application.name,
                        app_role="Viewer",
                        is_primary=False,
                        source="users_sheet",
                        remarks="Imported from Users Excel",
                    )
                )

    def _import_phase2(self) -> None:
        for row in self.phase2_rows:
            self.stats["phase2_rows_processed"] += 1
            app_name = normalize_whitespace(cell_text(row.get("app_name")))
            if app_name == "":
                self._skip(row, "missing App Names")
                continue

            key = normalize_application_name(app_name)
            existing = self.applications.get(key)
            if existing is not None:
                self.merged_applications.append(app_name)
                if existing.source == "phase2":
                    self.warnings.append(
                        f"PHASE2 row {row.get('_row')} duplicates application '{app_name}'. Merging."
                    )

            application = self._build_phase2_application(row, app_name, existing)
            self.applications[key] = application

            support_users = self._resolve_people(
                names=split_support_names(cell_text(row.get("support_name"))),
                emails=split_emails(cell_text(row.get("support_email"))),
                source="phase2_support",
                row_number=row.get("_row"),
                context=f"Support for {application.name}",
            )
            for person in support_users:
                self._add_assignment(
                    CanonicalAssignment(
                        user_email=person.email,
                        application_name=application.name,
                        app_role="Support",
                        is_primary=False,
                        source="phase2_support",
                        remarks="PHASE2 Support",
                    )
                )

            for person in self._resolve_people(
                names=split_support_names(cell_text(row.get("management_owner"))),
                emails=split_emails(cell_text(row.get("management_owner_email"))),
                source="phase2_management_owner",
                row_number=row.get("_row"),
                context=f"Management Owner for {application.name}",
            ):
                application.management_owner_emails.append(person.email)
                self._add_assignment(
                    CanonicalAssignment(
                        user_email=person.email,
                        application_name=application.name,
                        app_role="ZATCA Management",
                        is_primary=True,
                        source="phase2_management_owner",
                        remarks="PHASE2 ZATCA Management Owner",
                    )
                )

            for person in self._resolve_people(
                names=split_support_names(cell_text(row.get("app_lead"))),
                emails=split_emails(cell_text(row.get("app_lead_email"))),
                source="phase2_app_lead",
                row_number=row.get("_row"),
                context=f"App Lead for {application.name}",
            ):
                application.application_lead_emails.append(person.email)
                self._add_assignment(
                    CanonicalAssignment(
                        user_email=person.email,
                        application_name=application.name,
                        app_role="Application Lead",
                        is_primary=True,
                        source="phase2_app_lead",
                        remarks="PHASE2 ZATCA App Lead",
                    )
                )

            application.management_owner_emails = sorted(set(application.management_owner_emails))
            application.application_lead_emails = sorted(set(application.application_lead_emails))

    def _build_phase2_application(
        self,
        row: dict[str, Any],
        app_name: str,
        existing: CanonicalApplication | None,
    ) -> CanonicalApplication:
        type_raw = cell_text(row.get("application_type"))
        application_type = self._map_application_type(type_raw, app_name)
        department = self._map_department(type_raw, app_name)
        preferred_name = self.schema.preferred_application_name(app_name)
        if existing is not None:
            preferred_name = preferred_display_name(existing.name, preferred_name)

        technical_category = normalize_whitespace(cell_text(row.get("technical_category"))) or None
        description = normalize_whitespace(cell_text(row.get("description"))) or None
        vendor = normalize_whitespace(cell_text(row.get("vendor_name"))) or None
        remarks = normalize_whitespace(cell_text(row.get("remarks"))) or None
        is_live = parse_bool(row.get("live"))
        is_under_operation = parse_bool(row.get("under_operation"))
        live_raw = cell_text(row.get("live"))
        under_raw = cell_text(row.get("under_operation"))
        if live_raw and is_live is None:
            self.warnings.append(
                f"PHASE2 row {row.get('_row')} has non-boolean Live value '{live_raw}' for '{app_name}'."
            )
        if under_raw and is_under_operation is None:
            extra = normalize_whitespace(under_raw)
            remarks = _join_remarks(remarks, f"Under Operation: {extra}")
            self.warnings.append(
                f"PHASE2 row {row.get('_row')} has non-boolean under Operation value '{under_raw}' for '{app_name}'."
            )

        if vendor:
            vendor = self.schema.preferred_vendor_name(vendor)

        placeholder_fields: list[str] = list(existing.placeholder_fields) if existing else []
        if description is None:
            description = existing.description if existing and existing.description else PLACEHOLDER_REVIEW
            if "description" not in placeholder_fields:
                placeholder_fields.append("description")
                self._placeholder("Application", preferred_name, "description", description, "missing App Description")
        if technical_category is None:
            technical_category = (
                existing.technical_category
                if existing and existing.technical_category
                else PLACEHOLDER_TECHNICAL_CATEGORY
            )
            if "technical_category" not in placeholder_fields:
                placeholder_fields.append("technical_category")
                self._placeholder(
                    "Application",
                    preferred_name,
                    "technical_category",
                    technical_category,
                    "missing Technical Category",
                )
        if vendor is None:
            vendor = existing.vendor if existing and existing.vendor else PLACEHOLDER_VENDOR
            if "vendor" not in placeholder_fields:
                placeholder_fields.append("vendor")
                self._placeholder("Application", preferred_name, "vendor", vendor, "missing Vendor Name")

        technologies = split_technologies(cell_text(row.get("stack")))
        if existing:
            merged_tech: list[str] = []
            seen: set[str] = set()
            for name in existing.technologies + technologies:
                key = name.casefold()
                if key in seen:
                    continue
                seen.add(key)
                merged_tech.append(name)
            technologies = merged_tech
            remarks = remarks or existing.remarks

        return CanonicalApplication(
            name=preferred_name,
            application_type=application_type,
            department=department,
            technical_category=technical_category,
            description=description,
            vendor=vendor,
            support_type="Business Hours" if is_under_operation is not False else "Best Effort",
            status="Active" if is_live is not False else "Maintenance",
            criticality="Medium",
            technologies=technologies,
            remarks=remarks,
            source="phase2",
            management_owner_emails=list(existing.management_owner_emails) if existing else [],
            application_lead_emails=list(existing.application_lead_emails) if existing else [],
            placeholder_fields=placeholder_fields,
        )

    def _map_application_type(self, raw: str, app_name: str) -> str:
        text = normalize_whitespace(raw)
        if text == "":
            self._placeholder(
                "Application",
                app_name,
                "application_type",
                PLACEHOLDER_APPLICATION_TYPE,
                "missing Customs, Other Apps",
            )
            return PLACEHOLDER_APPLICATION_TYPE
        key = text.casefold()
        if key in APPLICATION_TYPE_ALIASES:
            mapped = APPLICATION_TYPE_ALIASES[key]
            return self.schema.preferred_application_type(mapped)
        matched = self.schema.preferred_application_type(text)
        known = {item.casefold() for item in self.schema.application_types}
        if matched.casefold() not in known:
            self.warnings.append(
                f"Application '{app_name}' uses unrecognised type '{text}'. "
                "The seeder will firstOrCreate it by name at runtime."
            )
        return matched

    def _map_department(self, type_raw: str, app_name: str) -> str:
        text = normalize_whitespace(type_raw)
        key = text.casefold()
        if key in APPLICATION_TYPE_ALIASES and APPLICATION_TYPE_ALIASES[key] == "Customs":
            return self.schema.preferred_department_name("Customs")
        if key in APPLICATION_TYPE_ALIASES and APPLICATION_TYPE_ALIASES[key] == "Internal App":
            return self.schema.preferred_department_name("Internal")
        if text == "":
            self._placeholder(
                "Application",
                app_name,
                "department",
                PLACEHOLDER_REVIEW,
                "missing Customs, Other Apps; department is required",
            )
            return self.schema.preferred_department_name(PLACEHOLDER_REVIEW)
        return self.schema.preferred_department_name(PLACEHOLDER_REVIEW)

    def _find_phase2_application(self, name: str) -> CanonicalApplication | None:
        key = normalize_application_name(name)
        found = self.applications.get(key)
        if found is not None:
            return found
        for application in self.applications.values():
            if normalize_application_name(application.name) == key:
                return application
        return None

    def _resolve_people(
        self,
        names: list[str],
        emails: list[str],
        source: str,
        row_number: Any,
        context: str,
    ) -> list[CanonicalUser]:
        if not names and not emails:
            return []
        if names and emails and len(names) == len(emails):
            pairs = list(zip(names, emails, strict=True))
        elif names:
            if emails and len(emails) != len(names):
                self.warnings.append(
                    f"PHASE2 row {row_number}: {context} name/email counts differ "
                    f"({len(names)} names, {len(emails)} emails); pairing by index and generating the rest."
                )
            pairs = []
            for index, name in enumerate(names):
                pairs.append((name, emails[index] if index < len(emails) else ""))
            for extra in emails[len(names) :]:
                pairs.append(("", extra))
        else:
            pairs = [("", email) for email in emails]

        people: list[CanonicalUser] = []
        for name, email in pairs:
            person = self._resolve_person(name, email, source, row_number, context)
            if person:
                people.append(person)
        return people

    def _resolve_person(
        self,
        name: str,
        email: str,
        source: str,
        row_number: Any,
        context: str,
    ) -> CanonicalUser | None:
        display = normalize_whitespace(name)
        explicit_email = normalize_email(email) if is_valid_email(normalize_email(email)) else ""
        if explicit_email == "" and email and not is_blank(email):
            self.warnings.append(f"PHASE2 row {row_number}: {context} has invalid email '{email}'.")
        if display == "" and explicit_email == "":
            return None
        if display == "" and explicit_email:
            local = explicit_email.split("@", 1)[0]
            display = normalize_whitespace(local.replace(".", " ").replace("_", " ")) or "Imported User"

        first, last = split_display_name(display)
        existing = self._find_user(display, explicit_email)
        if existing:
            if source not in existing.sources:
                existing.sources.append(source)
            if explicit_email and explicit_email != existing.email and explicit_email not in self.taken_emails:
                self.warnings.append(
                    f"PHASE2 row {row_number}: {context} matched existing user {existing.email} "
                    f"by name; Excel email '{explicit_email}' was not used to create a second account."
                )
            self.merged_users.append(existing.email)
            return existing

        generated = False
        resolved_email = explicit_email
        if not resolved_email:
            db_email = self.schema.user_email_by_name(display)
            if db_email:
                resolved_email = normalize_email(db_email)
            else:
                resolved_email, generated = self._make_generated_email(first, last, display)

        return self._upsert_user(
            CanonicalUser(
                first_name=first,
                last_name=last,
                email=resolved_email,
                display_name=display,
                email_source="phase2_explicit" if explicit_email else "generated",
                spatie_role="employee",
                sources=[source],
                generated_email=generated,
            )
        )

    def _find_user(self, display_name: str, email: str) -> CanonicalUser | None:
        if email:
            found = self.users.get(normalize_email(email))
            if found:
                return found
        name_key = normalize_name(display_name)
        if name_key and name_key in self.users_by_name:
            return self.users[self.users_by_name[name_key]]
        # Only reuse a generated address when it belongs to the same person.
        if display_name:
            first, last = split_display_name(display_name)
            generated = normalize_email(generate_email(first, last))
            found = self.users.get(generated)
            if found and found.generated_email and normalize_name(found.display_name) == name_key:
                return found
        return None

    def _upsert_user(self, incoming: CanonicalUser) -> CanonicalUser:
        incoming.email = normalize_email(incoming.email)
        existing = self._find_user(incoming.display_name, incoming.email)
        if existing is None:
            self.users[incoming.email] = incoming
            self.taken_emails.add(incoming.email)
            if incoming.display_name:
                self.users_by_name.setdefault(normalize_name(incoming.display_name), incoming.email)
            self.users_by_name.setdefault(
                normalize_name(f"{incoming.first_name} {incoming.last_name}"), incoming.email
            )
            return incoming

        if incoming.email != existing.email and incoming.email_source != "generated":
            pass
        for source in incoming.sources:
            if source not in existing.sources:
                existing.sources.append(source)
        if existing.phone is None and incoming.phone:
            existing.phone = incoming.phone
        if incoming.display_name:
            existing.display_name = preferred_display_name(existing.display_name, incoming.display_name)
            first, last = split_display_name(existing.display_name)
            existing.first_name = first
            existing.last_name = last
        if existing.static is False and incoming.spatie_role and existing.spatie_role == "employee":
            existing.spatie_role = incoming.spatie_role
        if incoming.static:
            existing.static = True
            existing.spatie_role = incoming.spatie_role
        self.users_by_name[normalize_name(existing.display_name or f"{existing.first_name} {existing.last_name}")] = (
            existing.email
        )
        return existing

    def _make_generated_email(self, first: str, last: str, display: str) -> tuple[str, bool]:
        base = generate_email(first, last)
        email, collided = unique_generated_email(base, self.taken_emails)
        self.generated_emails.append({"name": display or f"{first} {last}", "email": email})
        if collided:
            self.email_collisions.append({"name": display or f"{first} {last}", "from": base, "to": email})
        return email, True

    def _add_assignment(self, incoming: CanonicalAssignment) -> None:
        key = (normalize_email(incoming.user_email), normalize_application_name(incoming.application_name))
        existing = self.assignments.get(key)
        if existing is None:
            self.assignments[key] = incoming
            return
        current_rank = APP_ROLE_RANK.get(existing.app_role, 0)
        incoming_rank = APP_ROLE_RANK.get(incoming.app_role, 0)
        if incoming_rank > current_rank:
            existing.app_role = incoming.app_role
            existing.source = incoming.source
        if incoming.is_primary:
            existing.is_primary = True
        existing.remarks = _join_remarks(existing.remarks, incoming.remarks)

    def _placeholder(self, entity_type: str, entity: str, field: str, value: str, reason: str) -> None:
        self.placeholders.append(
            Placeholder(entity_type=entity_type, entity=entity, field=field, value=value, reason=reason)
        )

    def _skip(self, row: dict[str, Any], reason: str) -> None:
        self.skipped_rows.append(
            {"sheet": row.get("_sheet"), "row": row.get("_row"), "reason": reason}
        )
        self.warnings.append(
            f"{row.get('_sheet')} row {row.get('_row')} skipped: {reason}."
        )


def _join_remarks(current: str | None, incoming: str | None) -> str | None:
    left = (current or "").strip()
    right = (incoming or "").strip()
    if not left:
        return right or None
    if not right or right in left:
        return left
    return f"{left} | {right}"


# ---------------------------------------------------------------------------
# PHP seeder emitter
# ---------------------------------------------------------------------------


def emit_seeder(result: ImportResult, output_path: Path) -> Path:
    class_name = php_class_name(output_path)
    users_php = _php_array(
        [_user_payload(user) for user in result.users if not user.static],
        indent=8,
    )
    applications_php = _php_array([_application_payload(app) for app in result.applications], indent=8)
    assignments_php = _php_array([_assignment_payload(item) for item in result.assignments], indent=8)
    content = PHP_TEMPLATE.replace("__CLASS__", class_name)
    content = content.replace("__USERS__", users_php)
    content = content.replace("__APPLICATIONS__", applications_php)
    content = content.replace("__ASSIGNMENTS__", assignments_php)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(content, encoding="utf-8", newline="\n")
    return output_path


def _user_payload(user: CanonicalUser) -> dict[str, Any]:
    return {
        "first_name": user.first_name,
        "last_name": user.last_name,
        "email": user.email,
        "phone": user.phone,
        "role": user.spatie_role or "employee",
    }


def _application_payload(application: CanonicalApplication) -> dict[str, Any]:
    return {
        "name_en": application.name,
        "name_ar": application.name,
        "application_type": application.application_type,
        "department": application.department,
        "technical_category": application.technical_category,
        "description": application.description,
        "vendor": application.vendor,
        "status": application.status,
        "criticality": application.criticality,
        "support_type": application.support_type,
        "technologies": application.technologies,
        "remarks": application.remarks,
        "management_owner_emails": application.management_owner_emails,
        "application_lead_emails": application.application_lead_emails,
    }


def _assignment_payload(assignment: CanonicalAssignment) -> dict[str, Any]:
    return {
        "user_email": assignment.user_email,
        "application_name": assignment.application_name,
        "app_role": assignment.app_role,
        "is_primary": assignment.is_primary,
        "remarks": assignment.remarks,
    }


def _php_array(value: Any, indent: int = 0) -> str:
    spacer = " " * indent
    if isinstance(value, dict):
        if not value:
            return "[]"
        lines = ["["]
        for key, item in value.items():
            lines.append(f"{spacer}    {_php_string(str(key))} => {_php_array(item, indent + 4)},")
        lines.append(f"{spacer}]")
        return "\n".join(lines)
    if isinstance(value, list):
        if not value:
            return "[]"
        if all(not isinstance(item, (dict, list)) for item in value):
            return "[" + ", ".join(_php_array(item, 0) for item in value) + "]"
        lines = ["["]
        for item in value:
            lines.append(f"{spacer}    {_php_array(item, indent + 4)},")
        lines.append(f"{spacer}]")
        return "\n".join(lines)
    if isinstance(value, bool):
        return "true" if value else "false"
    if value is None:
        return "null"
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    return _php_string(str(value))


def _php_string(value: str) -> str:
    escaped = (
        value.replace("\\", "\\\\")
        .replace("'", "\\'")
        .replace("\r\n", "\n")
        .replace("\r", "\n")
        .replace("\n", "\\n")
    )
    return f"'{escaped}'"


PHP_TEMPLATE = r'''<?php

declare(strict_types=1);

namespace Database\Seeders;

use App\Enums\TechnologyCategory;
use App\Models\AppRole;
use App\Models\Application;
use App\Models\ApplicationAssignment;
use App\Models\ApplicationStatus;
use App\Models\ApplicationType;
use App\Models\Criticality;
use App\Models\Department;
use App\Models\SupportType;
use App\Models\Technology;
use App\Models\User;
use App\Models\Vendor;
use Illuminate\Database\Seeder;
use Illuminate\Support\Facades\DB;
use Illuminate\Support\Facades\Hash;
use Illuminate\Support\Facades\Schema;
use Illuminate\Support\Str;
use Spatie\Permission\Models\Role;

/**
 * Additive, idempotent import generated from Users + PHASE2 Excel files.
 * This seeder does not read Excel and never truncates existing data.
 */
class __CLASS__ extends Seeder
{
    private const string DEFAULT_PASSWORD = 'password';

    /**
     * Higher number wins when the same user is assigned to the same application
     * from more than one Excel source. The schema allows one OPEN assignment
     * per (application_id, user_id).
     *
     * @var array<string, int>
     */
    private const array APP_ROLE_RANK = [
        'ZATCA Management' => 40,
        'Application Lead' => 30,
        'Support' => 20,
        'Viewer' => 10,
    ];

    /** @var array<string, User> */
    private array $usersByNormalizedName = [];

    /** @var array<string, Application> */
    private array $applicationsByNormalizedName = [];

    public function run(): void
    {
        DB::transaction(function (): void {
            $this->ensureSystemRoles();
            $this->ensureAppRoles();

            $password = Hash::make(self::DEFAULT_PASSWORD);
            $assigner = $this->seedStaticUsers($password);
            $usersByEmail = $this->seedUsers($password);
            $applicationsByName = $this->seedApplications($assigner, $usersByEmail);
            $this->seedAssignments($assigner, $usersByEmail, $applicationsByName);
        });
    }

    private function ensureSystemRoles(): void
    {
        foreach (['super_admin', 'infra_admin', 'sd_admin', 'viewer', 'employee'] as $roleName) {
            Role::findOrCreate($roleName, 'web');
        }
    }

    private function ensureAppRoles(): void
    {
        $roles = [
            ['name' => 'Support', 'description' => 'Operational support access', 'sort_order' => 4],
            ['name' => 'ZATCA Management', 'description' => 'ZATCA management ownership and oversight', 'sort_order' => 5],
            ['name' => 'Viewer', 'description' => 'Read-only application access', 'sort_order' => 6],
            ['name' => 'Application Lead', 'description' => 'ZATCA application lead and technical ownership', 'sort_order' => 7],
        ];

        foreach ($roles as $role) {
            AppRole::query()->firstOrCreate(
                ['name' => $role['name']],
                [
                    'description' => $role['description'],
                    'is_active' => true,
                    'sort_order' => $role['sort_order'],
                ],
            );
        }
    }

    private function seedStaticUsers(string $password): User
    {
        $staticUsers = [
            [
                'first_name' => 'Super',
                'last_name' => 'Admin',
                'email' => 'super_admin@zatca.gov.sa',
                'phone' => '5678910110',
                'role' => 'super_admin',
            ],
            [
                'first_name' => 'Infra',
                'last_name' => 'Admin',
                'email' => 'infra_admin@zatca.gov.sa',
                'phone' => '5678910110',
                'role' => 'infra_admin',
            ],
            [
                'first_name' => 'SD',
                'last_name' => 'Admin',
                'email' => 'sd_admin@zatca.gov.sa',
                'phone' => '5678910110',
                'role' => 'sd_admin',
            ],
            [
                'first_name' => 'Viewer',
                'last_name' => 'User',
                'email' => 'viewer@zatca.gov.sa',
                'phone' => '5678910110',
                'role' => 'viewer',
            ],
        ];

        $superAdmin = null;

        foreach ($staticUsers as $row) {
            $user = $this->persistUser($row, $password);
            $this->assignSpatieRole($user, (string) $row['role']);
            if ($row['email'] === 'super_admin@zatca.gov.sa') {
                $superAdmin = $user;
            }
        }

        if ($superAdmin instanceof User) {
            return $superAdmin;
        }

        $existing = User::query()->role('super_admin')->orderBy('id')->first();
        if ($existing instanceof User) {
            return $existing;
        }

        return User::query()->firstOrCreate(
            ['email' => 'system@zatca.gov.sa'],
            [
                'first_name' => 'System',
                'last_name' => 'Importer',
                'password' => $password,
                'is_active' => true,
                'email_verified_at' => now(),
            ],
        );
    }

    /**
     * @return array<string, User>
     */
    private function seedUsers(string $password): array
    {
        $usersByEmail = [];

        foreach ($this->importedUsers() as $row) {
            $user = $this->persistUser($row, $password);
            $role = $row['role'] ?? 'employee';
            if (is_string($role) && $role !== '') {
                $this->assignSpatieRole($user, $role);
            }
            $usersByEmail[mb_strtolower((string) $user->email)] = $user;
        }

        foreach (User::withTrashed()->get() as $user) {
            $usersByEmail[mb_strtolower((string) $user->email)] = $user;
        }

        return $usersByEmail;
    }

    /**
     * @param  array<string, mixed>  $row
     */
    private function persistUser(array $row, string $password): User
    {
        $email = mb_strtolower(trim((string) $row['email']));
        $firstName = trim((string) $row['first_name']);
        $lastName = trim((string) $row['last_name']);
        $phone = $this->nullableString($row['phone'] ?? null);

        $user = User::withTrashed()->where('email', $email)->first();

        if (! $user instanceof User) {
            $this->bootUserNameIndex();
            $user = $this->usersByNormalizedName[$this->normalizePersonName($firstName.' '.$lastName)] ?? null;
        }

        if ($user instanceof User) {
            if ($user->trashed()) {
                $user->restore();
            }

            $updates = [];
            if ($user->first_name === '' || $user->first_name === null) {
                $updates['first_name'] = $firstName;
            }
            if ($user->last_name === '' || $user->last_name === null) {
                $updates['last_name'] = $lastName;
            }
            if (($user->phone === null || $user->phone === '') && $phone !== null) {
                $updates['phone'] = $phone;
            }
            if (! $user->is_active) {
                $updates['is_active'] = true;
            }
            if ($updates !== []) {
                $user->update($updates);
            }

            $this->rememberUser($user);

            return $user;
        }

        $created = User::query()->create([
            'first_name' => $firstName,
            'last_name' => $lastName,
            'email' => $email,
            'phone' => $phone,
            'password' => $password,
            'is_active' => true,
            'email_verified_at' => now(),
        ]);
        $this->rememberUser($created);

        return $created;
    }

    /**
     * @param  array<string, User>  $usersByEmail
     * @return array<string, Application>
     */
    private function seedApplications(User $assigner, array $usersByEmail): array
    {
        $applicationsByName = [];

        foreach ($this->importedApplications() as $row) {
            $name = trim((string) $row['name_en']);
            $normalized = $this->normalizeApplicationName($name);
            $application = $this->findApplication($name);

            $department = $this->firstOrCreateDepartment((string) $row['department']);
            $applicationType = $this->firstOrCreateApplicationType((string) $row['application_type']);
            $status = $this->findStatus((string) $row['status']);
            $criticality = $this->findCriticality((string) ($row['criticality'] ?? 'Medium'));
            $supportType = $this->findSupportType((string) $row['support_type']);
            $vendor = $this->firstOrCreateVendor($this->nullableString($row['vendor'] ?? null));

            if ($application instanceof Application) {
                if ($application->trashed()) {
                    $application->restore();
                }

                $payload = $this->applicationPayload(
                    $row,
                    $department->id,
                    $applicationType->id,
                    $status->id,
                    $criticality->id,
                    $supportType->id,
                    $vendor?->id,
                    $assigner->id,
                    includeCode: false,
                );
                $application->fill($this->withoutEmptyOverwrites($application, $payload));
                $application->save();
            } else {
                $payload = $this->applicationPayload(
                    $row,
                    $department->id,
                    $applicationType->id,
                    $status->id,
                    $criticality->id,
                    $supportType->id,
                    $vendor?->id,
                    $assigner->id,
                    includeCode: true,
                );
                $application = Application::query()->create($payload);
            }

            $this->rememberApplication($application);

            $technologyNames = $row['technologies'] ?? [];
            if (is_array($technologyNames) && $technologyNames !== []) {
                $technologyIds = [];
                foreach ($technologyNames as $technologyName) {
                    $technology = $this->firstOrCreateTechnology((string) $technologyName);
                    $technologyIds[] = $technology->id;
                }
                $application->technologies()->syncWithoutDetaching($technologyIds);
            }

            $ownerEmails = $row['management_owner_emails'] ?? [];
            $leadEmails = $row['application_lead_emails'] ?? [];
            if (is_array($ownerEmails) && $ownerEmails !== []) {
                $application->businessOwners()->syncWithoutDetaching(
                    $this->userIdsFromEmails($ownerEmails, $usersByEmail)
                );
            }
            if (is_array($leadEmails) && $leadEmails !== []) {
                $application->technicalOwners()->syncWithoutDetaching(
                    $this->userIdsFromEmails($leadEmails, $usersByEmail)
                );
            }

            $applicationsByName[$normalized] = $application;
        }

        foreach (Application::withTrashed()->get() as $application) {
            $applicationsByName[$this->normalizeApplicationName((string) $application->name_en)] = $application;
        }

        return $applicationsByName;
    }

    /**
     * @param  array<string, mixed>  $row
     * @return array<string, mixed>
     */
    private function applicationPayload(
        array $row,
        int $departmentId,
        int $applicationTypeId,
        int $statusId,
        int $criticalityId,
        int $supportTypeId,
        ?int $vendorId,
        int $actorId,
        bool $includeCode,
    ): array {
        $payload = [
            'name_en' => trim((string) $row['name_en']),
            'name_ar' => trim((string) $row['name_ar']),
            'department_id' => $departmentId,
            'application_type_id' => $applicationTypeId,
            'status_id' => $statusId,
            'criticality_id' => $criticalityId,
            'support_type_id' => $supportTypeId,
            'updated_by' => $actorId,
        ];

        if ($includeCode) {
            $payload['code'] = $this->uniqueApplicationCode((string) $row['name_en']);
            $payload['created_by'] = $actorId;
        }

        if (Schema::hasColumn('applications', 'description')) {
            $payload['description'] = $this->nullableString($row['description'] ?? null);
        }
        if (Schema::hasColumn('applications', 'technical_category')) {
            $payload['technical_category'] = $this->nullableString($row['technical_category'] ?? null);
        }
        if (Schema::hasColumn('applications', 'vendor_id')) {
            $payload['vendor_id'] = $vendorId;
        }
        if (Schema::hasColumn('applications', 'remarks')) {
            $payload['remarks'] = $this->nullableString($row['remarks'] ?? null);
        }

        return $payload;
    }

    /**
     * @param  array<string, mixed>  $payload
     * @return array<string, mixed>
     */
    private function withoutEmptyOverwrites(Application $application, array $payload): array
    {
        $filtered = [];

        foreach ($payload as $key => $value) {
            if ($value === null || $value === '') {
                continue;
            }

            $current = $application->getAttribute($key);
            if (
                is_string($current)
                && trim($current) !== ''
                && in_array($key, ['description', 'remarks', 'technical_category'], true)
                && is_string($value)
                && (str_starts_with($value, 'TODO') || str_starts_with($value, 'Unknown'))
            ) {
                continue;
            }

            $filtered[$key] = $value;
        }

        return $filtered;
    }

    /**
     * @param  array<string, User>  $usersByEmail
     * @param  array<string, Application>  $applicationsByName
     */
    private function seedAssignments(User $assigner, array $usersByEmail, array $applicationsByName): void
    {
        $roles = AppRole::query()->get()->keyBy(fn (AppRole $role): string => mb_strtolower($role->name));

        foreach ($this->importedAssignments() as $row) {
            $user = $usersByEmail[mb_strtolower((string) $row['user_email'])] ?? null;
            $application = $applicationsByName[$this->normalizeApplicationName((string) $row['application_name'])] ?? null;
            $role = $roles->get(mb_strtolower((string) $row['app_role']));

            if (! $user instanceof User || ! $application instanceof Application || ! $role instanceof AppRole) {
                continue;
            }

            $existing = ApplicationAssignment::query()
                ->where('application_id', $application->id)
                ->where('user_id', $user->id)
                ->open()
                ->first();

            $incomingRank = self::APP_ROLE_RANK[$role->name] ?? 0;
            $remarks = $this->nullableString($row['remarks'] ?? null);
            $isPrimary = (bool) ($row['is_primary'] ?? false);

            if ($existing instanceof ApplicationAssignment) {
                $currentRoleName = (string) ($existing->appRole()->value('name') ?? '');
                $currentRank = self::APP_ROLE_RANK[$currentRoleName] ?? 0;
                $updates = [];
                if ($incomingRank > $currentRank) {
                    $updates['app_role_id'] = $role->id;
                }
                if ($isPrimary && ! $existing->is_primary) {
                    $updates['is_primary'] = true;
                }
                if ($remarks !== null) {
                    $currentRemarks = trim((string) ($existing->remarks ?? ''));
                    if ($currentRemarks === '') {
                        $updates['remarks'] = $remarks;
                    } elseif (! str_contains($currentRemarks, $remarks)) {
                        $updates['remarks'] = $currentRemarks.' | '.$remarks;
                    }
                }
                if ($updates !== []) {
                    $existing->update($updates);
                }
                continue;
            }

            ApplicationAssignment::query()->create([
                'application_id' => $application->id,
                'user_id' => $user->id,
                'app_role_id' => $role->id,
                'assigned_by' => $assigner->id,
                'assigned_at' => now(),
                'ended_at' => null,
                'is_primary' => $isPrimary,
                'remarks' => $remarks,
            ]);
        }
    }

    private function findApplication(string $name): ?Application
    {
        $this->bootApplicationNameIndex();

        return $this->applicationsByNormalizedName[$this->normalizeApplicationName($name)] ?? null;
    }

    private function bootUserNameIndex(): void
    {
        if ($this->usersByNormalizedName !== []) {
            return;
        }

        foreach (User::withTrashed()->get() as $user) {
            $this->rememberUser($user);
        }
    }

    private function rememberUser(User $user): void
    {
        $this->usersByNormalizedName[$this->normalizePersonName(
            trim((string) $user->first_name.' '.(string) $user->last_name)
        )] = $user;
    }

    private function bootApplicationNameIndex(): void
    {
        if ($this->applicationsByNormalizedName !== []) {
            return;
        }

        foreach (Application::withTrashed()->get() as $application) {
            $this->rememberApplication($application);
        }
    }

    private function rememberApplication(Application $application): void
    {
        $this->applicationsByNormalizedName[$this->normalizeApplicationName((string) $application->name_en)] = $application;
        $this->applicationsByNormalizedName[$this->normalizeApplicationName((string) $application->name_ar)] = $application;
    }

    private function firstOrCreateDepartment(string $name): Department
    {
        $name = trim($name) !== '' ? trim($name) : 'TODO - Needs Review';
        $normalized = $this->normalizeApplicationName($name);
        $existing = Department::withTrashed()->get()->first(
            function (Department $department) use ($normalized): bool {
                $existingName = $this->normalizeApplicationName((string) $department->name_en);
                if ($existingName === $normalized) {
                    return true;
                }
                if ($normalized === 'customs' && str_contains($existingName, 'custom')) {
                    return true;
                }
                if ($normalized === 'internal' && (str_contains($existingName, 'zakat') || str_contains($existingName, 'tax'))) {
                    return true;
                }

                return false;
            }
        );

        if ($existing instanceof Department) {
            if ($existing->trashed()) {
                $existing->restore();
            }

            return $existing;
        }

        return Department::query()->firstOrCreate(
            ['name_en' => $name],
            ['name_ar' => $name],
        );
    }

    private function firstOrCreateApplicationType(string $name): ApplicationType
    {
        $name = trim($name) !== '' ? trim($name) : 'Internal App';
        $normalized = $this->normalizeApplicationName($name);
        $existing = ApplicationType::query()->get()->first(
            fn (ApplicationType $type): bool => $this->normalizeApplicationName((string) $type->name_en) === $normalized
                || $this->normalizeApplicationName((string) $type->code) === $normalized
        );

        if ($existing instanceof ApplicationType) {
            return $existing;
        }

        $code = Str::upper(Str::slug($name, '_'));
        if ($code === '') {
            $code = 'APP_TYPE';
        }

        return ApplicationType::query()->firstOrCreate(
            ['code' => $code],
            [
                'name_en' => $name,
                'name_ar' => $name,
            ],
        );
    }

    private function firstOrCreateVendor(?string $name): ?Vendor
    {
        if ($name === null || trim($name) === '') {
            return null;
        }

        $name = trim($name);
        $normalized = $this->normalizeApplicationName($name);
        $existing = Vendor::withTrashed()->get()->first(
            fn (Vendor $vendor): bool => $this->normalizeApplicationName((string) $vendor->name) === $normalized
        );

        if ($existing instanceof Vendor) {
            if ($existing->trashed()) {
                $existing->restore();
            }

            return $existing;
        }

        return Vendor::query()->firstOrCreate(
            ['name' => $name],
            [
                'status' => true,
                'remarks' => str_starts_with($name, 'TODO') || str_starts_with($name, 'Unknown')
                    ? 'Placeholder vendor created by Excel import'
                    : 'Imported from PHASE2',
            ],
        );
    }

    private function firstOrCreateTechnology(string $name): Technology
    {
        $name = trim($name);
        $normalized = mb_strtolower($name);
        $existing = Technology::withTrashed()->get()->first(
            fn (Technology $technology): bool => mb_strtolower((string) $technology->name) === $normalized
        );

        if ($existing instanceof Technology) {
            if ($existing->trashed()) {
                $existing->restore();
            }

            return $existing;
        }

        return Technology::query()->firstOrCreate(
            ['name' => $name],
            [
                'category' => TechnologyCategory::Other,
                'description' => 'Imported from PHASE2 stack and technologies',
                'is_active' => true,
            ],
        );
    }

    private function findStatus(string $code): ApplicationStatus
    {
        $status = ApplicationStatus::query()->where('code', $code)->first()
            ?? ApplicationStatus::query()->where('code', 'Active')->first();

        if ($status instanceof ApplicationStatus) {
            return $status;
        }

        return ApplicationStatus::query()->firstOrCreate(
            ['code' => 'Active'],
            ['name_en' => 'Active', 'name_ar' => 'نشط', 'is_active' => true],
        );
    }

    private function findCriticality(string $code): Criticality
    {
        $criticality = Criticality::query()->where('code', $code)->first()
            ?? Criticality::query()->where('code', 'Medium')->first();

        if ($criticality instanceof Criticality) {
            return $criticality;
        }

        return Criticality::query()->firstOrCreate(
            ['code' => 'Medium'],
            ['name_en' => 'Medium', 'name_ar' => 'متوسط', 'is_active' => true],
        );
    }

    private function findSupportType(string $code): SupportType
    {
        $supportType = SupportType::query()->where('code', $code)->first()
            ?? SupportType::query()->where('code', 'Business Hours')->first();

        if ($supportType instanceof SupportType) {
            return $supportType;
        }

        return SupportType::query()->firstOrCreate(
            ['code' => 'Business Hours'],
            ['name_en' => 'Business Hours', 'name_ar' => 'ساعات العمل', 'is_active' => true],
        );
    }

    private function uniqueApplicationCode(string $appName): string
    {
        $base = Str::upper(Str::slug($appName, '_'));
        if ($base === '') {
            $base = 'APP';
        }
        $base = Str::limit($base, 90, '');
        $code = $base;
        $suffix = 1;

        while (Application::withTrashed()->where('code', $code)->exists()) {
            $code = Str::limit($base, 80, '').'_'.$suffix;
            $suffix++;
        }

        return $code;
    }

    /**
     * @param  list<string>  $emails
     * @param  array<string, User>  $usersByEmail
     * @return list<int>
     */
    private function userIdsFromEmails(array $emails, array $usersByEmail): array
    {
        $ids = [];
        foreach ($emails as $email) {
            $user = $usersByEmail[mb_strtolower(trim((string) $email))] ?? null;
            if ($user instanceof User) {
                $ids[] = (int) $user->id;
            }
        }

        return array_values(array_unique($ids));
    }

    private function assignSpatieRole(User $user, string $roleName): void
    {
        if (! $user->hasRole($roleName)) {
            $user->assignRole($roleName);
        }
    }

    private function normalizePersonName(string $name): string
    {
        return mb_strtolower(trim((string) preg_replace('/\s+/u', ' ', $name)));
    }

    private function normalizeApplicationName(string $name): string
    {
        return mb_strtolower(trim((string) preg_replace('/\s+/u', ' ', $name)));
    }

    private function nullableString(mixed $value): ?string
    {
        if (! is_string($value)) {
            return null;
        }

        $trimmed = trim($value);

        return $trimmed === '' ? null : $trimmed;
    }

    /**
     * @return list<array<string, mixed>>
     */
    private function importedUsers(): array
    {
        return __USERS__;
    }

    /**
     * @return list<array<string, mixed>>
     */
    private function importedApplications(): array
    {
        return __APPLICATIONS__;
    }

    /**
     * @return list<array<string, mixed>>
     */
    private function importedAssignments(): array
    {
        return __ASSIGNMENTS__;
    }
}
'''


# ---------------------------------------------------------------------------
# Report
# ---------------------------------------------------------------------------


def write_report(
    result: ImportResult,
    report_path: Path,
    seeder_path: Path,
    users_file: Path,
    applications_file: Path,
    users_detection: SheetDetection,
    phase2_detection: SheetDetection,
    schema: SchemaSnapshot,
) -> Path:
    lines = [
        "Excel to Laravel seeder — migration report",
        f"Generated at: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}",
        f"Users file: {users_file}",
        f"Applications file: {applications_file}",
        f"Users sheet: {users_detection.sheet_name} (header row {users_detection.header_row})",
        f"PHASE2 sheet: {phase2_detection.sheet_name} (header row {phase2_detection.header_row})",
        f"Seeder: {seeder_path}",
        f"MySQL connected: {'yes' if schema.connected else 'no'}"
        + (f" ({schema.database})" if schema.database else ""),
        "",
        "== Summary ==",
        f"Users rows processed: {result.stats['users_rows_processed']}",
        f"PHASE2 rows processed: {result.stats['phase2_rows_processed']}",
        f"Unique users: {result.stats['unique_users']}",
        f"Users merged: {result.stats['users_merged']}",
        f"Applications: {result.stats['applications']}",
        f"Assignments: {result.stats['assignments']}",
        f"Generated emails: {result.stats['generated_emails']}",
        f"Email collisions: {result.stats['email_collisions']}",
        f"Placeholder values: {result.stats['placeholder_values']}",
        f"Invalid/skipped rows: {result.stats['invalid_skipped_rows']}",
        f"Unmatched user applications (user created, no assignment): {result.stats['unmatched_user_applications']}",
        "",
    ]

    lines.append("== Static users always included ==")
    for row in STATIC_USERS:
        lines.append(f"  {row['email']}  role={row['role']}  phone={row['phone']}")
    lines.append("")

    lines.append("== Generated emails ==")
    if result.generated_emails:
        for item in result.generated_emails:
            lines.append(f"  {item['name']} -> {item['email']}")
    else:
        lines.append("  (none)")
    lines.append("")

    lines.append("== Email collisions ==")
    if result.email_collisions:
        for item in result.email_collisions:
            lines.append(f"  {item['name']}: {item['from']} -> {item['to']}")
    else:
        lines.append("  (none)")
    lines.append("")

    lines.append("== Users merged ==")
    if result.merged_users:
        for email in result.merged_users:
            lines.append(f"  {email}")
    else:
        lines.append("  (none)")
    lines.append("")

    lines.append("== Applications merged ==")
    if result.merged_applications:
        for name in result.merged_applications:
            lines.append(f"  {name}")
    else:
        lines.append("  (none)")
    lines.append("")

    lines.append("== Unmatched user applications ==")
    lines.append("Users were still created. Assignments were skipped because the app is not in PHASE2.")
    if result.unmatched_user_applications:
        for item in result.unmatched_user_applications:
            lines.append(
                f"  {item.get('sheet')} row {item.get('row')}: {item.get('user')} -> {item.get('application')}"
            )
    else:
        lines.append("  (none)")
    lines.append("")

    lines.append("== Placeholder values ==")
    if result.placeholders:
        for item in result.placeholders:
            lines.append(
                f"  [{item.entity_type}] {item.entity} / {item.field} = {item.value} ({item.reason})"
            )
    else:
        lines.append("  (none)")
    lines.append("")

    lines.append("== Invalid/skipped rows ==")
    if result.skipped_rows:
        for item in result.skipped_rows:
            lines.append(f"  {item.get('sheet')} row {item.get('row')}: {item.get('reason')}")
    else:
        lines.append("  (none)")
    lines.append("")

    lines.append("== Warnings ==")
    if result.warnings or schema.warnings:
        for warning in schema.warnings + result.warnings:
            lines.append(f"  {warning}")
    else:
        lines.append("  (none)")
    lines.append("")

    if schema.connected:
        lines.append("== Live schema snapshot (read-only) ==")
        lines.append(f"  application types: {', '.join(schema.application_types) or '(none)'}")
        lines.append(f"  departments: {', '.join(schema.departments) or '(none)'}")
        lines.append(f"  app roles: {', '.join(schema.app_roles) or '(none)'}")
        lines.append(f"  spatie roles: {', '.join(schema.roles) or '(none)'}")
        lines.append("")

    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")

    sidecar = report_path.with_suffix(".json")
    sidecar.write_text(
        json.dumps(
            {
                "stats": result.stats,
                "generated_emails": result.generated_emails,
                "email_collisions": result.email_collisions,
                "merged_users": result.merged_users,
                "unmatched_user_applications": result.unmatched_user_applications,
                "placeholders": [item.__dict__ for item in result.placeholders],
                "skipped_rows": result.skipped_rows,
                "warnings": schema.warnings + result.warnings,
            },
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
        newline="\n",
    )
    return report_path


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Read Users + PHASE2 Excel files and generate a self-contained Laravel seeder.",
    )
    parser.add_argument("users_xlsx", type=Path, help="Path to the Users Excel file")
    parser.add_argument("applications_xlsx", type=Path, help="Path to the Applications Excel file (PHASE2 sheet)")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("ImportedPortfolioSeeder.php"),
        help="Output PHP seeder file (or directory). Default: ImportedPortfolioSeeder.php",
    )
    parser.add_argument(
        "--report",
        type=Path,
        default=None,
        help="Output migration report path. Default: migration_report.txt next to the seeder",
    )
    parser.add_argument("--db-host", default=DEFAULT_DB["host"])
    parser.add_argument("--db-port", type=int, default=int(DEFAULT_DB["port"]))
    parser.add_argument("--db-database", default=DEFAULT_DB["database"])
    parser.add_argument("--db-username", default=DEFAULT_DB["user"])
    parser.add_argument("--db-password", default=DEFAULT_DB["password"])
    parser.add_argument(
        "--skip-db",
        action="store_true",
        help="Skip live MySQL inspection and use schema fallbacks",
    )
    return parser.parse_args(argv)


def resolve_output_path(output: Path) -> Path:
    path = output.expanduser()
    if path.exists() and path.is_dir():
        return path / "ImportedPortfolioSeeder.php"
    if path.suffix.lower() != ".php":
        return path / "ImportedPortfolioSeeder.php"
    return path


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    users_file = args.users_xlsx.expanduser().resolve()
    applications_file = args.applications_xlsx.expanduser().resolve()

    if not users_file.is_file():
        raise SystemExit(f"Users Excel file not found: {users_file}")
    if not applications_file.is_file():
        raise SystemExit(f"Applications Excel file not found: {applications_file}")

    output_path = resolve_output_path(args.output).resolve()
    report_path = (
        args.report.expanduser().resolve()
        if args.report is not None
        else output_path.with_name("migration_report.txt")
    )

    print(f"Inspecting MySQL schema (read-only): {args.db_database}@{args.db_host}:{args.db_port}")
    schema = inspect_schema(
        host=args.db_host,
        port=args.db_port,
        database=args.db_database,
        user=args.db_username,
        password=args.db_password,
        skip_db=args.skip_db,
    )
    if schema.connected:
        print(
            f"Connected. tables={len(schema.tables)} users={len(schema.users)} "
            f"applications={len(schema.applications)} types={len(schema.application_types)}"
        )
    else:
        print("MySQL was not inspected; seeder will resolve related records by name at runtime.")
        for warning in schema.warnings:
            print(f"Schema warning: {warning}")

    print(f"Reading users file: {users_file}")
    users_rows, users_detection = load_users_rows(users_file)
    print(
        f"Users sheet '{users_detection.sheet_name}' header row {users_detection.header_row}: "
        f"{len(users_rows)} data rows"
    )

    print(f"Reading applications file: {applications_file}")
    phase2_rows, phase2_detection, sheet_names = load_phase2_rows(applications_file)
    print(f"Workbook sheets: {', '.join(sheet_names)}")
    print(
        f"PHASE2 sheet '{phase2_detection.sheet_name}' header row {phase2_detection.header_row}: "
        f"{len(phase2_rows)} data rows"
    )

    result = ImportPipeline(users_rows, phase2_rows, schema).run()
    seeder_path = emit_seeder(result, output_path)
    report = write_report(
        result,
        report_path,
        seeder_path,
        users_file,
        applications_file,
        users_detection,
        phase2_detection,
        schema,
    )

    print(f"Wrote seeder: {seeder_path}")
    print(f"Wrote report: {report}")
    print(
        "Summary: "
        f"users_rows={result.stats['users_rows_processed']} "
        f"phase2_rows={result.stats['phase2_rows_processed']} "
        f"unique_users={result.stats['unique_users']} "
        f"merged={result.stats['users_merged']} "
        f"apps={result.stats['applications']} "
        f"assignments={result.stats['assignments']} "
        f"generated_emails={result.stats['generated_emails']} "
        f"collisions={result.stats['email_collisions']} "
        f"placeholders={result.stats['placeholder_values']} "
        f"skipped={result.stats['invalid_skipped_rows']} "
        f"unmatched_user_apps={result.stats['unmatched_user_applications']}"
    )
    print("Next: copy the seeder into backend/database/seeders and run:")
    print(f"  php artisan db:seed --class={php_class_name(seeder_path)}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except HeaderMismatchError as exc:
        print(str(exc), file=sys.stderr)
        raise SystemExit(2) from exc
